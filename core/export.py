import re
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from django.utils import timezone

from core import ui

CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
EXPORT_LIMIT = 10000
HEADER_ROW = 5
WINE = "920000"


def sheet_title(text):
    return re.sub(r"[\\/?*\[\]:]", " ", text).strip()[:31] or "Datos"


def file_name(slug):
    return f"{slug}-{timezone.localdate().isoformat()}.xlsx"


def as_excel_value(value):
    if isinstance(value, datetime):
        if timezone.is_aware(value):
            value = timezone.localtime(value).replace(tzinfo=None)
        return value
    if isinstance(value, Decimal):
        return float(value)
    return value


def build_workbook(title, headers, kinds, rows, summary, user, total):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title(title)
    workbook.properties.title = title
    workbook.properties.creator = "SGR"

    sheet["A1"] = title
    sheet["A1"].font = Font(name="Calibri", size=14, bold=True, color=WINE)
    stamp = timezone.localtime().strftime("%d/%m/%Y %H:%M")
    sheet["A2"] = f"Generado el {stamp} por {user.get_username()}"
    sheet["A3"] = summary
    if total > len(rows):
        sheet["A4"] = (
            f"Se exportaron las primeras {ui.fmt_number(len(rows))} filas de "
            f"{ui.fmt_number(total)}. Agrega filtros para acotar el resultado."
        )
        sheet["A4"].font = Font(italic=True, color="92191A")
    for row in (2, 3):
        sheet[f"A{row}"].font = Font(color="47433E")

    header_fill = PatternFill("solid", fgColor=WINE)
    thin = Side(style="thin", color="DAD7D1")
    for index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=HEADER_ROW, column=index, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    widths = [len(str(header)) + 2 for header in headers]
    for offset, values in enumerate(rows, start=1):
        for index, (value, kind) in enumerate(zip(values, kinds, strict=True), start=1):
            cell = sheet.cell(row=HEADER_ROW + offset, column=index)
            value = as_excel_value(value)
            if isinstance(value, str):
                cell.value = value
                cell.data_type = "s"
            else:
                cell.value = value
            if isinstance(value, datetime):
                cell.number_format = "DD/MM/YYYY HH:MM"
            elif isinstance(value, date):
                cell.number_format = "DD/MM/YYYY"
            elif isinstance(value, (int, float)) and not isinstance(value, bool):
                cell.number_format = '0.0"%"' if kind == "percent" else "#,##0.##"
                cell.alignment = Alignment(horizontal="right")
            cell.border = Border(bottom=thin)
            widths[index - 1] = max(widths[index - 1], min(len(str(value)) + 2, 60))

    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(10, min(width, 60))

    last_row = HEADER_ROW + max(len(rows), 1)
    sheet.freeze_panes = f"A{HEADER_ROW + 1}"
    sheet.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(len(headers))}{last_row}"
    sheet.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
