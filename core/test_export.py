"""Pruebas de la descarga a Excel: contenido, filtros, límite, permisos y auditoría."""

from datetime import datetime
from io import BytesIO
from types import SimpleNamespace
from unittest import mock
from zoneinfo import ZoneInfo

from django.conf import settings
from django.test import SimpleTestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from core import export
from core.crud import PAGE_SIZE, registry
from core.models import AuditLog
from core.testing import PortalTestCase
from organization.models import Beneficiary, Officer


class WorkbookTests(SimpleTestCase):
    """El libro que arma core/export.py, sin pasar por las pantallas."""

    def build(self, rows, total=None, **changes):
        arguments = {
            "title": "Prueba",
            "headers": ["Nombre", "Cantidad", "Fecha"],
            "kinds": ["text", "number", "date"],
            "rows": rows,
            "summary": "Sin filtros",
            "user": SimpleNamespace(get_username=lambda: "ana"),
            "total": len(rows) if total is None else total,
        }
        arguments.update(changes)
        return load_workbook(BytesIO(export.build_workbook(**arguments))).active

    def test_sheet_names_are_valid_for_excel(self):
        self.assertEqual(export.sheet_title("Informe: de [gestión]/2026?"), "Informe  de  gestión  2026")
        self.assertEqual(len(export.sheet_title("x" * 100)), 31)
        self.assertEqual(export.sheet_title("///"), "Datos")

    def test_file_name_has_the_day(self):
        self.assertEqual(export.file_name("actividades"), f"actividades-{timezone.localdate().isoformat()}.xlsx")

    def test_header_rows_notes_and_table(self):
        sheet = self.build([["Ana", 3, datetime(2026, 10, 1, 9, 30)]], summary="Búsqueda: «ana»")
        self.assertEqual(sheet.title, "Prueba")
        self.assertEqual(sheet["A1"].value, "Prueba")
        self.assertTrue(sheet["A2"].value.startswith("Generado el "))
        self.assertTrue(sheet["A2"].value.endswith(" por ana"))
        self.assertEqual(sheet["A3"].value, "Búsqueda: «ana»")
        self.assertIsNone(sheet["A4"].value)
        self.assertEqual([c.value for c in sheet[export.HEADER_ROW]], ["Nombre", "Cantidad", "Fecha"])
        self.assertEqual(sheet.freeze_panes, "A6")
        self.assertEqual(sheet.auto_filter.ref, "A5:C6")

    def test_numbers_and_dates_are_typed_cells(self):
        sheet = self.build([["Ana", 3, datetime(2026, 10, 1, 9, 30)], ["Luis", 2.5, ""]])
        self.assertEqual(sheet["B6"].value, 3)
        self.assertEqual(sheet["B7"].value, 2.5)
        self.assertEqual(sheet["C6"].value, datetime(2026, 10, 1, 9, 30))
        self.assertEqual(sheet["C6"].number_format, "DD/MM/YYYY HH:MM")

    def test_percent_columns_keep_the_number_and_show_a_percent_sign(self):
        sheet = self.build([["Ana", 52.5, ""]], kinds=["text", "percent", "date"])
        self.assertEqual(sheet["B6"].value, 52.5)
        self.assertEqual(sheet["B6"].number_format, '0.0"%"')

    def test_texts_that_look_like_formulas_stay_text(self):
        sheet = self.build([["=1+1", 1, ""], ["+56911111111", 1, ""], ["@SUM(A1)", 1, ""]])
        self.assertEqual(sheet["A6"].value, "=1+1")
        self.assertEqual(sheet["A6"].data_type, "s")
        self.assertEqual(sheet["A7"].value, "+56911111111")
        self.assertEqual(sheet["A8"].value, "@SUM(A1)")
        self.assertEqual(sheet["A8"].data_type, "s")

    def test_a_cut_result_says_how_many_rows_were_left_out(self):
        sheet = self.build([["Ana", 1, ""], ["Luis", 2, ""]], total=1200)
        self.assertEqual(
            sheet["A4"].value,
            "Se exportaron las primeras 2 filas de 1.200. Agrega filtros para acotar el resultado.",
        )

    def test_an_empty_result_still_has_its_header(self):
        sheet = self.build([])
        self.assertEqual([c.value for c in sheet[export.HEADER_ROW]], ["Nombre", "Cantidad", "Fecha"])
        self.assertEqual(sheet.auto_filter.ref, "A5:C6")


class ExcelTestCase(PortalTestCase):
    def download(self, url_name, **params):
        response = self.client.get(reverse(url_name), {"formato": "xlsx", **params})
        self.assertEqual(response.status_code, 200)
        return response

    def sheet(self, response):
        return load_workbook(BytesIO(response.content)).active

    def table(self, sheet):
        """Encabezados y filas de datos de la hoja."""
        rows = [list(row) for row in sheet.iter_rows(min_row=export.HEADER_ROW, values_only=True)]
        return rows[0], rows[1:]

    def column(self, response, header):
        headers, rows = self.table(self.sheet(response))
        return [row[headers.index(header)] for row in rows]


class ListExportTests(ExcelTestCase):
    def test_download_is_an_excel_file_named_after_the_list_and_the_day(self):
        self.login_as("admin")
        response = self.download("core:actividades_list")
        self.assertEqual(response["Content-Type"], export.CONTENT_TYPE)
        # El nombre lleva el día en la hora de Chile (la que usan las pantallas), no el día en UTC: de noche
        # en Chile ya es "mañana" en UTC y la prueba fallaba sin que hubiera ningún error.
        with timezone.override(ZoneInfo(settings.PORTAL_TIME_ZONE)):
            today = timezone.localdate().isoformat()
        self.assertEqual(response["Content-Disposition"], f'attachment; filename="actividades-{today}.xlsx"')

    def test_sheet_has_title_author_filters_header_and_rows(self):
        self.login_as("admin")
        crud = registry.get("actividades")
        sheet = self.sheet(self.download("core:actividades_list"))
        self.assertEqual(sheet["A1"].value, crud.list_title)
        self.assertTrue(sheet["A2"].value.endswith(" por admin"))
        self.assertEqual(sheet["A3"].value, "Sin filtros")
        headers, rows = self.table(sheet)
        self.assertEqual(headers, [column.label for column in crud.columns])
        self.assertEqual(len(rows), 5)  # las 6 de demostración menos la archivada
        self.assertEqual(sheet.auto_filter.ref, f"A5:{get_column_letter(len(headers))}10")

    def test_download_follows_the_search_the_filters_and_the_order_of_the_screen(self):
        self.login_as("admin")
        response = self.download("core:actividades_list", q="munoz")
        self.assertEqual(self.column(response, "Contacto"), ["Carla Muñoz"])
        self.assertIn("Búsqueda: «munoz»", self.sheet(response)["A3"].value)

        response = self.download("core:actividades_list", status="closed")
        self.assertEqual(self.column(response, "Contacto"), ["Pedro Díaz"])
        self.assertIn("Cerrada", self.sheet(response)["A3"].value)

        response = self.download("core:actividades_list", orden="contact_name", dir="desc")
        self.assertEqual(
            self.column(response, "Contacto"),
            ["Pedro Díaz", "María Pérez", "Luis Herrera", "Carla Muñoz", "Ana Rojas"],
        )

        response = self.download("core:actividades_list", archivados="1")
        self.assertEqual(len(self.column(response, "Contacto")), 6)
        self.assertIn("incluye archivados", self.sheet(response)["A3"].value)

    def test_download_ignores_the_page_and_exports_all_the_rows(self):
        Beneficiary.objects.bulk_create(
            [Beneficiary(rut=f"9{number:07d}-1", name=f"Persona {number}") for number in range(45)]
        )
        self.login_as("admin")
        names = self.column(self.download("core:beneficiarios_list", page=2), "Nombre")
        self.assertEqual(len(names), Beneficiary.objects.count())
        self.assertGreater(len(names), PAGE_SIZE)

    def test_download_is_cut_at_the_limit_and_says_so(self):
        self.login_as("admin")
        with mock.patch.object(export, "EXPORT_LIMIT", 3):
            sheet = self.sheet(self.download("core:actividades_list"))
        _, rows = self.table(sheet)
        self.assertEqual(len(rows), 3)
        self.assertEqual(
            sheet["A4"].value,
            "Se exportaron las primeras 3 filas de 5. Agrega filtros para acotar el resultado.",
        )
        # Sin recorte no hay aviso.
        self.assertIsNone(self.sheet(self.download("core:actividades_list"))["A4"].value)

    def test_a_delegation_user_only_gets_its_own_rows(self):
        self.login_as("matias.cavieres")
        response = self.download("core:actividades_list")
        self.assertEqual(set(self.column(response, "Contacto")), {"María Pérez", "Carla Muñoz"})
        self.assertEqual(set(self.column(response, "Delegación")), {"Delegación Norte"})

    def test_download_needs_the_same_permission_as_the_list(self):
        self.login_as("jefatura1")
        response = self.client.get(reverse("core:usuarios_list"), {"formato": "xlsx"})
        self.assertEqual(response.status_code, 403)
        self.assertFalse(AuditLog.objects.filter(event="Listado exportado a Excel").exists())

    def test_anonymous_visitors_are_sent_to_the_login(self):
        response = self.client.get(reverse("core:actividades_list"), {"formato": "xlsx"})
        self.assertEqual(response.status_code, 302)
        self.assertTrue(response["Location"].startswith(reverse("core:login") + "?next="))

    def test_every_download_is_audited_and_a_plain_visit_is_not(self):
        user = self.login_as("admin")
        self.client.get(reverse("core:actividades_list"), {"q": "munoz"})
        self.assertFalse(AuditLog.objects.filter(event="Listado exportado a Excel").exists())
        self.download("core:actividades_list", q="munoz")
        log = AuditLog.objects.get(event="Listado exportado a Excel")
        self.assertEqual(log.user, user)
        self.assertEqual(log.affected_entity, "Activity")
        self.assertEqual(log.record_id, 0)
        self.assertTrue(log.new_value.startswith("1 fila. "))
        self.assertIn("Búsqueda: «munoz»", log.new_value)
        self.download("core:actividades_list")
        self.assertIn("5 filas. Sin filtros", AuditLog.objects.order_by("-pk").first().new_value)

    def test_texts_that_look_like_formulas_stay_text(self):
        Beneficiary.objects.create(rut="90000001-1", name="=1+1", address="@SUM(A1)")
        self.login_as("admin")
        sheet = self.sheet(self.download("core:beneficiarios_list"))
        found = {
            cell.value: cell.data_type
            for row in sheet.iter_rows(min_row=export.HEADER_ROW + 1)
            for cell in row
            if cell.value in ("=1+1", "@SUM(A1)")
        }
        self.assertEqual(found, {"=1+1": "s", "@SUM(A1)": "s"})

    def test_numbers_and_dates_keep_their_type(self):
        self.login_as("admin")
        response = self.download("core:indicadores_list")
        headers, rows = self.table(self.sheet(response))
        self.assertTrue(rows)
        for row in rows:
            self.assertIsInstance(row[headers.index("Avance")], (int, float))
            self.assertIsInstance(row[headers.index("Fecha de cálculo")], datetime)
        percent = self.sheet(response).cell(
            row=export.HEADER_ROW + 1, column=headers.index("Cumplimiento") + 1
        )
        self.assertEqual(percent.number_format, '0.0"%"')

    def test_every_list_can_be_downloaded(self):
        """Todos los listados que se ven en el menú tienen su botón de Excel funcionando."""
        self.login_as("admin")
        for crud in registry.all():
            with self.subTest(crud=crud.slug):
                response = self.client.get(crud.url("list"), {"formato": "xlsx"})
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response["Content-Type"], export.CONTENT_TYPE)
                load_workbook(BytesIO(response.content))  # el archivo se puede abrir


class ReportExportTests(ExcelTestCase):
    def test_report_workbook_has_one_row_per_officer(self):
        self.login_as("coordinacion1")
        response = self.client.get(reverse("core:reports"), {"formato": "xlsx"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], export.CONTENT_TYPE)
        self.assertIn('filename="informe-sgr-', response["Content-Disposition"])
        sheet = self.sheet(response)
        headers, rows = self.table(sheet)
        self.assertEqual(sheet["A1"].value, "Informe de gestión")
        self.assertEqual(headers[:3], ["Funcionario", "Delegación", "Cargo"])
        self.assertEqual(len(rows), Officer.objects.filter(status="active", deleted_at__isnull=True).count())
        quiroz = next(row for row in rows if row[0] == "Matías Quiroz")
        self.assertEqual(quiroz[headers.index("Cumplimiento promedio (%)")], 52.5)

    def test_report_download_is_audited(self):
        self.login_as("coordinacion1")
        self.client.get(reverse("core:reports"), {"formato": "xlsx"})
        log = AuditLog.objects.get(event="Listado exportado a Excel")
        self.assertEqual(log.affected_entity, "Informe")

    def test_a_delegation_user_only_gets_its_own_officers(self):
        self.login_as("jefatura1")
        response = self.client.get(reverse("core:reports"), {"formato": "xlsx"})
        self.assertEqual(response.status_code, 200)
        headers, rows = self.table(self.sheet(response))
        self.assertEqual({row[headers.index("Delegación")] for row in rows}, {"Delegación Norte"})

    def test_the_report_page_links_to_the_excel_file(self):
        self.login_as("coordinacion1")
        response = self.client.get(reverse("core:reports"))
        self.assertContains(response, "formato=xlsx")
        self.assertContains(response, "Excel")
