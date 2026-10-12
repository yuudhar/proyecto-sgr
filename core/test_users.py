"""Pruebas de Control de usuarios: quién entra, crear, editar, contraseñas, perfiles y funcionario."""

from io import BytesIO

from django.contrib.auth.models import Group, User
from django.urls import NoReverseMatch, reverse
from openpyxl import load_workbook

from accounts.models import UserProfile
from core import export
from core.models import AuditLog
from core.testing import PASSWORD, PortalTestCase
from organization.models import Officer

NEW_PASSWORD = "ClaveSegura#2026"


class UserTestCase(PortalTestCase):
    def group(self, name):
        return Group.objects.get(name=name).pk

    def user_data(self, **changes):
        """Datos del formulario de un usuario nuevo y válido (ana.torres, de la delegación Norte)."""
        data = {
            "username": "ana.torres",
            "first_name": "Ana",
            "last_name": "Torres",
            "email": "ana.torres@dominio.cl",
            "groups": [self.group("Funcionario")],
            "officer": self.officer("F-1004").pk,
            "is_active": "on",
            "password1": NEW_PASSWORD,
            "password2": NEW_PASSWORD,
        }
        data.update(changes)
        return {key: value for key, value in data.items() if value is not None}

    def edit_data(self, user, **changes):
        """Datos de un usuario que ya existe, como los envía el formulario al guardar sin tocar nada."""
        profile = getattr(user, "profile", None)
        data = {
            "username": user.username,
            "first_name": user.first_name,
            "last_name": user.last_name,
            "email": user.email,
            "groups": [group.pk for group in user.groups.all()],
            "officer": profile.officer_id if profile is not None and profile.officer_id else "",
            "is_active": "on" if user.is_active else None,
            "is_staff": "on" if user.is_staff else None,
            "is_superuser": "on" if user.is_superuser else None,
            "password1": "",
            "password2": "",
        }
        data.update(changes)
        return {key: value for key, value in data.items() if value is not None}

    def update_url(self, user):
        return reverse("core:usuarios_update", args=[user.pk])

    def create(self, **changes):
        return self.client.post(reverse("core:usuarios_create"), self.user_data(**changes))


class UserAccessTests(UserTestCase):
    def test_only_the_superuser_opens_the_user_screens(self):
        for username in ("jefatura1", "matias.cavieres", "coordinacion1"):
            with self.subTest(username=username):
                self.login_as(username)
                self.assertEqual(self.client.get(reverse("core:usuarios_list")).status_code, 403)
                self.assertEqual(self.client.get(reverse("core:usuarios_create")).status_code, 403)
                self.assertEqual(self.create().status_code, 403)
        self.assertFalse(User.objects.filter(username="ana.torres").exists())
        self.login_as("admin")
        self.assertEqual(self.client.get(reverse("core:usuarios_list")).status_code, 200)
        self.assertEqual(self.client.get(reverse("core:usuarios_create")).status_code, 200)

    def test_the_menu_entry_is_only_for_the_superuser(self):
        self.login_as("admin")
        self.assertIn("Control de usuarios", self.menu_labels(self.client.get(reverse("core:dashboard"))))
        self.login_as("coordinacion1")
        self.assertNotIn("Control de usuarios", self.menu_labels(self.client.get(reverse("core:dashboard"))))

    def test_a_non_superuser_cannot_edit_or_see_another_account(self):
        quiroz = User.objects.get(username="matias.quiroz")
        self.login_as("matias.cavieres")
        self.assertEqual(self.client.get(reverse("core:usuarios_detail", args=[quiroz.pk])).status_code, 403)
        before = quiroz.email
        response = self.client.post(self.update_url(quiroz), self.edit_data(quiroz, email="x@dominio.cl"))
        self.assertEqual(response.status_code, 403)
        quiroz.refresh_from_db()
        self.assertEqual(quiroz.email, before)

    def test_users_are_deactivated_not_archived(self):
        with self.assertRaises(NoReverseMatch):
            reverse("core:usuarios_archive", args=[1])
        self.login_as("admin")
        quiroz = User.objects.get(username="matias.quiroz")
        self.assertNotContains(self.client.get(reverse("core:usuarios_detail", args=[quiroz.pk])), "Eliminar")
        self.assertNotContains(self.client.get(self.update_url(quiroz)), "Eliminar")


class UserListTests(UserTestCase):
    def cells(self, response):
        return {row.obj.username: [cell.text for cell in row.cells] for row in response.context["rows"]}

    def test_the_list_shows_name_profiles_and_delegation_of_each_user(self):
        self.login_as("admin")
        cells = self.cells(self.client.get(reverse("core:usuarios_list")))
        self.assertEqual(
            set(cells),
            {"admin", "consulta1", "coordinacion1", "jefatura1", "matias.cavieres", "matias.quiroz", "supervision1"},
        )
        # Columnas: usuario, nombre, correo, perfiles, delegación, activo, último ingreso
        self.assertEqual(cells["admin"][1:5], ["", "admin@sgr.example", "Super Usuario", "Todas las delegaciones"])
        self.assertEqual(
            cells["coordinacion1"][1:5],
            ["", "coordinacion1@sgr.example", "Coordinación", "Todas las delegaciones"],
        )
        self.assertEqual(
            cells["jefatura1"][1:5], ["Jefatura Norte", "jefatura1@sgr.example", "Jefatura", "Delegación Norte"]
        )
        self.assertEqual(
            cells["matias.cavieres"][1:5],
            ["Matías Cavieres", "matias.cavieres@sgr.example", "Funcionario, Verificador", "Delegación Norte"],
        )
        self.assertEqual(cells["matias.quiroz"][4], "Delegación Sur")
        self.assertEqual(
            cells["consulta1"][1:5],
            ["Consulta Norte", "consulta1@sgr.example", "Consulta", "Delegación Norte"],
        )
        # Con delegaciones adicionales (A2) la columna dice cuántas otras puede elegir.
        self.assertEqual(
            cells["supervision1"][1:5],
            ["Supervisión Zonal", "supervision1@sgr.example", "Jefatura", "Delegación Norte y 1 más"],
        )

    def test_filters_by_profile_and_by_state(self):
        self.login_as("admin")
        url = reverse("core:usuarios_list")
        response = self.client.get(url, {"groups": self.group("Funcionario")})
        self.assertEqual(set(self.cells(response)), {"matias.cavieres", "matias.quiroz"})
        response = self.client.get(url, {"groups": self.group("Jefatura")})
        self.assertEqual(set(self.cells(response)), {"jefatura1", "supervision1"})
        self.assertEqual(self.cells(self.client.get(url, {"is_active": "0"})), {})
        self.assertEqual(len(self.cells(self.client.get(url, {"is_active": "1"}))), 7)
        response = self.client.get(url, {"groups": self.group("Consulta")})
        self.assertEqual(set(self.cells(response)), {"consulta1"})

    def test_every_sortable_column_works(self):
        self.login_as("admin")
        response = self.client.get(reverse("core:usuarios_list"))
        keys = [column.sort_key for column in response.context["columns"] if column.sort_key]
        self.assertTrue({"username", "first_name", "profile__officer__delegation__name"} <= set(keys))
        for key in keys:
            for direction in ("asc", "desc"):
                with self.subTest(key=key, direction=direction):
                    response = self.client.get(reverse("core:usuarios_list"), {"orden": key, "dir": direction})
                    self.assertEqual(response.status_code, 200)
                    self.assertEqual(len(response.context["rows"]), 7)
        response = self.client.get(reverse("core:usuarios_list"), {"orden": "username", "dir": "desc"})
        self.assertEqual(response.context["rows"][0].obj.username, "supervision1")  # la última por orden alfabético

    def test_the_excel_file_has_the_same_columns_and_never_the_password(self):
        self.login_as("admin")
        response = self.client.get(reverse("core:usuarios_list"), {"formato": "xlsx"})
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.content)).active
        headers = [cell.value for cell in sheet[export.HEADER_ROW]]
        self.assertEqual(headers, ["Usuario", "Nombre", "Correo", "Perfiles", "Delegación", "Activo", "Último ingreso"])
        values = {
            cell.value for row in sheet.iter_rows(min_row=export.HEADER_ROW + 1) for cell in row if cell.value
        }
        self.assertIn("Funcionario, Verificador", values)
        self.assertFalse([value for value in values if "pbkdf2" in str(value)])

    def test_the_detail_shows_profiles_and_officer_and_hides_the_password(self):
        self.login_as("admin")
        cavieres = User.objects.get(username="matias.cavieres")
        response = self.client.get(reverse("core:usuarios_detail", args=[cavieres.pk]))
        self.assertEqual(response.status_code, 200)
        for text in ("Funcionario, Verificador", "Matías Cavieres", "Delegación Norte", "Encargado Social", "Cuenta creada el"):
            self.assertContains(response, text)
        self.assertNotContains(response, "pbkdf2")


class UserCreateTests(UserTestCase):
    def test_the_form_offers_the_profiles_and_only_free_active_officers(self):
        self.login_as("admin")
        response = self.client.get(reverse("core:usuarios_create"))
        self.assertContains(response, 'class="role-name"', count=5)  # una casilla por perfil, con su descripción
        for name in ("Coordinación", "Jefatura", "Funcionario", "Verificador", "Consulta"):
            self.assertContains(response, name)
        offered = {officer.institutional_id for officer in response.context["form"].fields["officer"].queryset}
        self.assertEqual(offered, {"F-1002", "F-1004", "F-2002", "F-2004", "F-2005"})  # sin los que ya tienen usuario
        Officer.objects.filter(institutional_id="F-1004").update(status=Officer.Status.INACTIVE)
        response = self.client.get(reverse("core:usuarios_create"))
        offered = {officer.institutional_id for officer in response.context["form"].fields["officer"].queryset}
        self.assertNotIn("F-1004", offered)

    def test_the_form_lists_what_a_good_password_needs(self):
        self.login_as("admin")
        response = self.client.get(reverse("core:usuarios_create"))
        self.assertContains(response, "Debe tener al menos 8 caracteres.")
        self.assertNotContains(response, "clave")

    def test_creating_a_user_with_profile_officer_and_password(self):
        admin = self.login_as("admin")
        response = self.create()
        self.assertEqual(response.status_code, 302)
        user = User.objects.get(username="ana.torres")
        self.assertRedirects(response, reverse("core:usuarios_detail", args=[user.pk]))
        self.assertTrue(user.check_password(NEW_PASSWORD))
        self.assertEqual([group.name for group in user.groups.all()], ["Funcionario"])
        self.assertTrue(user.is_active)
        self.assertFalse(user.is_staff)
        self.assertFalse(user.is_superuser)
        self.assertEqual(user.get_full_name(), "Ana Torres")
        self.assertEqual(user.profile.officer, self.officer("F-1004"))
        self.assertEqual(user.profile.status, UserProfile.Status.ACTIVE)
        log = AuditLog.objects.get(event="Registro creado", affected_entity="User", record_id=user.pk)
        self.assertEqual(log.user, admin)
        self.assertEqual(log.new_value, "ana.torres")  # el nombre de la cuenta y nada más

    def test_the_new_user_can_log_in_and_only_sees_its_delegation(self):
        self.login_as("admin")
        self.create()
        self.client.logout()
        self.assertTrue(self.client.login(username="ana.torres", password=NEW_PASSWORD))
        response = self.client.get(reverse("core:actividades_list"))
        self.assertEqual({row.obj.delegation_id for row in response.context["rows"]}, {self.north.pk})
        self.assertEqual(self.client.get(reverse("core:usuarios_list")).status_code, 403)

    def test_save_and_add_another_goes_back_to_an_empty_form(self):
        self.login_as("admin")
        response = self.create(_addanother="1")
        self.assertRedirects(response, reverse("core:usuarios_create"))
        self.assertTrue(User.objects.filter(username="ana.torres").exists())

    def test_the_account_follows_the_format_rules(self):
        self.login_as("admin")
        response = self.create(username="a b", first_name="Ana2", email="no-es-correo")
        self.assertEqual(response.status_code, 200)
        for name in ("username", "first_name", "email"):
            self.assertIn(name, response.context["form"].errors)
        self.assertFalse(User.objects.filter(username="a b").exists())

    def test_the_username_cannot_repeat_even_with_other_capitals(self):
        self.login_as("admin")
        response = self.create(username="MATIAS.CAVIERES")
        self.assertEqual(response.status_code, 200)
        self.assertIn("username", response.context["form"].errors)

    def test_the_email_is_required_because_the_password_recovery_depends_on_it(self):
        self.login_as("admin")
        response = self.create(email="")
        self.assertEqual(response.status_code, 200)
        self.assertIn("email", response.context["form"].errors)
        self.assertFalse(User.objects.filter(username="ana.torres").exists())
        self.assertContains(self.client.get(reverse("core:usuarios_create")), "No puede repetirse entre usuarios.")

    def test_the_email_cannot_repeat_even_with_other_capitals(self):
        self.login_as("admin")
        response = self.create(email="MATIAS.CAVIERES@SGR.EXAMPLE")  # ya es de matias.cavieres
        self.assertEqual(response.status_code, 200)
        self.assertIn("Ese correo ya lo usa otro usuario", " ".join(response.context["form"].errors["email"]))
        self.assertFalse(User.objects.filter(username="ana.torres").exists())

    def test_passwords_are_required_equal_and_not_weak(self):
        self.login_as("admin")
        response = self.create(password1="", password2="")
        self.assertIn("password1", response.context["form"].errors)
        self.assertIn("password2", response.context["form"].errors)

        response = self.create(password2="otra distinta")
        self.assertEqual(response.context["form"].errors["password2"], ["Las dos contraseñas deben ser iguales."])

        response = self.create(password1="12345678", password2="12345678")
        message = " ".join(response.context["form"].errors["password1"])
        self.assertIn("contraseña", message)
        self.assertFalse(User.objects.filter(username="ana.torres").exists())

        response = self.create(password1="ana.torres2026", password2="ana.torres2026")  # se parece al nombre de usuario
        self.assertIn("password1", response.context["form"].errors)
        self.assertFalse(User.objects.filter(username="ana.torres").exists())

    def test_a_user_needs_a_profile_or_to_be_superuser(self):
        self.login_as("admin")
        response = self.create(groups=[])
        self.assertIn("groups", response.context["form"].errors)
        response = self.create(groups=[], officer="", is_superuser="on", is_staff="on")
        self.assertEqual(response.status_code, 302)
        user = User.objects.get(username="ana.torres")
        self.assertTrue(user.is_superuser)
        self.assertEqual(user.profile.officer, None)

    def test_the_officer_is_required_unless_superuser_or_coordination(self):
        self.login_as("admin")
        response = self.create(officer="")
        self.assertIn("officer", response.context["form"].errors)
        response = self.create(officer="", groups=[self.group("Coordinación")])
        self.assertEqual(response.status_code, 302)
        user = User.objects.get(username="ana.torres")
        self.assertIsNone(user.profile.officer)
        self.assertEqual([group.name for group in user.groups.all()], ["Coordinación"])

    def test_an_officer_can_only_belong_to_one_user(self):
        self.login_as("admin")
        response = self.create(officer=self.officer("F-1001").pk)  # ya es de matias.cavieres
        self.assertEqual(response.status_code, 200)
        self.assertIn("officer", response.context["form"].errors)
        self.assertFalse(User.objects.filter(username="ana.torres").exists())

    def test_a_forged_group_that_is_not_a_profile_choice_is_rejected(self):
        self.login_as("admin")
        response = self.create(groups=[99999])
        self.assertEqual(response.status_code, 200)
        self.assertIn("groups", response.context["form"].errors)


class UserEditTests(UserTestCase):
    def setUp(self):
        super().setUp()
        self.admin = self.login_as("admin")
        self.quiroz = User.objects.get(username="matias.quiroz")

    def test_editing_keeps_the_password_when_it_is_left_blank(self):
        response = self.client.post(self.update_url(self.quiroz), self.edit_data(self.quiroz, email="quiroz@dominio.cl"))
        self.assertEqual(response.status_code, 302)
        self.quiroz.refresh_from_db()
        self.assertEqual(self.quiroz.email, "quiroz@dominio.cl")
        self.assertTrue(self.quiroz.check_password(PASSWORD))
        log = AuditLog.objects.get(event="Registro modificado", affected_entity="User", record_id=self.quiroz.pk)
        self.assertIn("quiroz@dominio.cl", log.new_value)
        self.assertNotIn("pbkdf2", f"{log.previous_value} {log.new_value}")

    def test_an_account_cannot_take_the_email_of_another(self):
        data = self.edit_data(self.quiroz, email="Matias.Cavieres@sgr.example")
        response = self.client.post(self.update_url(self.quiroz), data)
        self.assertEqual(response.status_code, 200)
        self.assertIn("email", response.context["form"].errors)
        self.quiroz.refresh_from_db()
        self.assertEqual(self.quiroz.email, "matias.quiroz@sgr.example")

    def test_an_old_account_without_email_can_still_be_edited(self):
        # Una cuenta antigua sin correo no se bloquea: se puede desactivar o cambiar de perfil sin inventarle uno.
        User.objects.filter(pk=self.quiroz.pk).update(email="")
        self.quiroz.refresh_from_db()
        form = self.client.get(self.update_url(self.quiroz)).context["form"]
        self.assertFalse(form.fields["email"].required)
        response = self.client.post(self.update_url(self.quiroz), self.edit_data(self.quiroz, first_name="Matías"))
        self.assertEqual(response.status_code, 302)
        self.quiroz.refresh_from_db()
        self.assertEqual((self.quiroz.first_name, self.quiroz.email), ("Matías", ""))

    def test_an_account_that_has_an_email_cannot_lose_it(self):
        response = self.client.post(self.update_url(self.quiroz), self.edit_data(self.quiroz, email=""))
        self.assertEqual(response.status_code, 200)
        self.assertIn("email", response.context["form"].errors)

    def test_a_repeated_email_does_not_block_editing_other_data_of_an_old_account(self):
        # Dos cuentas antiguas con el mismo correo: solo se exige que sea distinto cuando se cambia el correo.
        User.objects.filter(pk=self.quiroz.pk).update(email="matias.cavieres@sgr.example")
        self.quiroz.refresh_from_db()
        response = self.client.post(self.update_url(self.quiroz), self.edit_data(self.quiroz, first_name="Matías"))
        self.assertEqual(response.status_code, 302)
        response = self.client.post(
            self.update_url(self.quiroz), self.edit_data(self.quiroz, email="Matias.Cavieres@SGR.example")
        )
        self.assertEqual(response.status_code, 302)  # sigue siendo el mismo correo (solo cambian las mayúsculas)

    def test_saving_without_changes_leaves_no_audit_entry(self):
        response = self.client.post(self.update_url(self.quiroz), self.edit_data(self.quiroz))
        self.assertEqual(response.status_code, 302)
        self.assertFalse(AuditLog.objects.filter(affected_entity="User", record_id=self.quiroz.pk).exists())

    def test_the_edit_form_shows_the_current_profiles_and_officer(self):
        response = self.client.get(self.update_url(self.quiroz))
        form = response.context["form"]
        selected = {int(pk) for pk in form["groups"].value()}
        self.assertEqual(selected, {self.group("Funcionario"), self.group("Verificador")})
        self.assertEqual(form["officer"].value(), self.officer("F-2001").pk)
        self.assertContains(response, "Déjala en blanco para conservar la contraseña actual.")
        self.assertNotContains(response, "pbkdf2")

    def test_changing_profiles_is_audited(self):
        data = self.edit_data(self.quiroz, groups=[self.group("Funcionario")])
        self.assertEqual(self.client.post(self.update_url(self.quiroz), data).status_code, 302)
        self.assertEqual([group.name for group in self.quiroz.groups.all()], ["Funcionario"])
        log = AuditLog.objects.get(event="Perfil o funcionario modificado", record_id=self.quiroz.pk)
        self.assertIn("Perfiles: Funcionario, Verificador", log.previous_value)
        self.assertIn("Perfiles: Funcionario;", log.new_value)

    def test_changing_the_officer_moves_the_user_to_another_delegation(self):
        data = self.edit_data(self.quiroz, officer=self.officer("F-1004").pk)
        self.assertEqual(self.client.post(self.update_url(self.quiroz), data).status_code, 302)
        self.quiroz.refresh_from_db()
        self.assertEqual(self.quiroz.profile.officer, self.officer("F-1004"))
        log = AuditLog.objects.get(event="Perfil o funcionario modificado", record_id=self.quiroz.pk)
        self.assertIn("Matías Quiroz", log.previous_value)
        self.client.logout()
        self.assertTrue(self.client.login(username="matias.quiroz", password=PASSWORD))
        response = self.client.get(reverse("core:actividades_list"))
        self.assertEqual({row.obj.delegation_id for row in response.context["rows"]}, {self.north.pk})

    def test_the_superuser_can_reset_a_password(self):
        data = self.edit_data(self.quiroz, password1="NuevaClave#2026", password2="NuevaClave#2026")
        self.assertEqual(self.client.post(self.update_url(self.quiroz), data).status_code, 302)
        self.quiroz.refresh_from_db()
        self.assertTrue(self.quiroz.check_password("NuevaClave#2026"))
        log = AuditLog.objects.get(event="Contraseña restablecida", record_id=self.quiroz.pk)
        self.assertEqual(log.user, self.admin)
        self.assertFalse(log.new_value)  # la contraseña no queda anotada
        self.assertFalse(AuditLog.objects.filter(new_value__contains="NuevaClave").exists())
        self.assertFalse(AuditLog.objects.filter(new_value__contains="pbkdf2").exists())

    def test_a_weak_or_mismatched_reset_changes_nothing(self):
        for first, second in (("12345678", "12345678"), ("NuevaClave#2026", "otra distinta")):
            with self.subTest(first=first, second=second):
                data = self.edit_data(self.quiroz, password1=first, password2=second)
                self.assertEqual(self.client.post(self.update_url(self.quiroz), data).status_code, 200)
                self.quiroz.refresh_from_db()
                self.assertTrue(self.quiroz.check_password(PASSWORD))

    def test_the_superuser_cannot_deactivate_or_demote_their_own_account(self):
        response = self.client.post(self.update_url(self.admin), self.edit_data(self.admin, is_active=None))
        self.assertIn("is_active", response.context["form"].errors)
        response = self.client.post(self.update_url(self.admin), self.edit_data(self.admin, is_superuser=None))
        self.assertIn("is_superuser", response.context["form"].errors)
        self.admin.refresh_from_db()
        self.assertTrue(self.admin.is_active)
        self.assertTrue(self.admin.is_superuser)

    def test_deactivating_another_user_blocks_the_login_and_marks_the_profile(self):
        data = self.edit_data(self.quiroz, is_active=None)
        self.assertEqual(self.client.post(self.update_url(self.quiroz), data).status_code, 302)
        self.quiroz.refresh_from_db()
        self.assertFalse(self.quiroz.is_active)
        self.assertEqual(self.quiroz.profile.status, UserProfile.Status.INACTIVE)
        self.client.logout()
        self.assertFalse(self.client.login(username="matias.quiroz", password=PASSWORD))

        self.login_as("admin")
        self.client.post(self.update_url(self.quiroz), self.edit_data(self.quiroz, is_active="on"))
        self.quiroz.refresh_from_db()
        self.assertTrue(self.quiroz.is_active)
        self.assertEqual(self.quiroz.profile.status, UserProfile.Status.ACTIVE)

    def test_a_global_user_without_officer_can_be_edited(self):
        coordination = User.objects.get(username="coordinacion1")
        response = self.client.post(self.update_url(coordination), self.edit_data(coordination, email="coord@dominio.cl"))
        self.assertEqual(response.status_code, 302)
        coordination.refresh_from_db()
        self.assertEqual(coordination.email, "coord@dominio.cl")
