"""Pruebas del espacio de trabajo: menú lateral con delegación y perfiles, Atención, Mi perfil y Administración."""

from io import BytesIO

from django.contrib.auth.models import User
from django.test import SimpleTestCase
from django.urls import reverse
from openpyxl import load_workbook

from core import export
from core.colors import DEFAULT_COLOR, clean_color, contrast, palette
from core.models import AuditLog
from core.testing import PASSWORD, PortalTestCase
from organization.models import Delegation


class ColorTests(SimpleTestCase):
    def test_a_color_is_written_in_lowercase_or_replaced_by_the_default(self):
        self.assertEqual(clean_color(" #1D4E89 "), "#1d4e89")
        for bad in ("", None, "rojo", "#12345", "#12345G", "red;}", "#1D4E89;x"):
            self.assertEqual(clean_color(bad), DEFAULT_COLOR, repr(bad))

    def test_dark_backgrounds_get_white_text_and_light_ones_black(self):
        self.assertEqual(palette("#1D4E89")["fg"], "#ffffff")
        self.assertEqual(palette("#920000")["fg"], "#ffffff")
        self.assertEqual(palette("#F5E6A0")["fg"], "#000000")

    def test_the_text_is_readable_over_any_color(self):
        """Contraste de al menos 4,5 (WCAG AA) para el texto, también sobre el sombreado de las etiquetas."""
        steps = range(0, 256, 51)
        for red in steps:
            for green in steps:
                for blue in steps:
                    color = f"#{red:02x}{green:02x}{blue:02x}"
                    colors = palette(color)
                    self.assertGreaterEqual(contrast(colors["bg"], colors["fg"]), 4.5, color)
                    self.assertGreaterEqual(contrast(self.shaded(colors), colors["fg"]), 4.5, color)

    @staticmethod
    def shaded(colors):
        """El fondo de una etiqueta: el color de la delegación con el sombreado encima."""
        numbers = colors["shade"][5:-1].split(",")  # "rgba(0, 0, 0, 0.26)"
        overlay = [int(value) for value in numbers[:3]]
        alpha = float(numbers[3])
        base = [int(colors["bg"][index : index + 2], 16) for index in (1, 3, 5)]
        mixed = [round(b * (1 - alpha) + o * alpha) for b, o in zip(base, overlay)]
        return "#" + "".join(f"{value:02x}" for value in mixed)


class SidebarTests(PortalTestCase):
    def page(self, username, name="core:dashboard"):
        self.login_as(username)
        response = self.client.get(reverse(name))
        self.assertEqual(response.status_code, 200)
        return response

    def test_each_user_sees_its_delegation_and_its_profiles(self):
        cases = {
            "admin": ("Todas las delegaciones", ["Super Usuario"]),
            "coordinacion1": ("Todas las delegaciones", ["Coordinación"]),
            "jefatura1": ("Delegación Norte", ["Jefatura"]),
            "matias.cavieres": ("Delegación Norte", ["Funcionario", "Verificador"]),
            "matias.quiroz": ("Delegación Sur", ["Funcionario", "Verificador"]),
        }
        for username, (scope, roles) in cases.items():
            with self.subTest(username=username):
                response = self.page(username)
                self.assertEqual(response.context["me_scope"]["name"], scope)
                self.assertEqual([role["label"] for role in response.context["me_roles"]], roles)
                self.assertContains(response, f'class="scope-name">{scope}</p>')
                for role in roles:
                    self.assertContains(response, role)

    def test_the_block_says_profile_or_profiles_by_how_many_there_are(self):
        self.assertContains(self.page("jefatura1"), 'id="perfiles-titulo">Perfil</p>')
        self.assertContains(self.page("matias.cavieres"), 'id="perfiles-titulo">Perfiles</p>')

    def test_each_delegation_has_its_own_color(self):
        self.assertContains(self.page("matias.cavieres"), "--scope-bg: #1d4e89")
        self.assertContains(self.page("matias.quiroz"), "--scope-bg: #2e7d5b")
        self.assertContains(self.page("admin"), f"--scope-bg: {DEFAULT_COLOR}")

    def test_a_color_changed_in_the_maintainer_shows_up_in_the_menu(self):
        Delegation.objects.filter(pk=self.north.pk).update(color="#112233")
        response = self.page("matias.cavieres")
        self.assertContains(response, "--scope-bg: #112233")
        self.assertContains(response, "--scope-fg: #ffffff")
        Delegation.objects.filter(pk=self.north.pk).update(color="#F5E6A0")
        response = self.page("matias.cavieres")
        self.assertContains(response, "--scope-bg: #f5e6a0")
        self.assertContains(response, "--scope-fg: #000000")

    def test_a_stored_color_that_is_not_a_color_never_reaches_the_page(self):
        Delegation.objects.filter(pk=self.north.pk).update(color="red;}</style><b>x")
        response = self.page("matias.cavieres")
        self.assertNotContains(response, "</style><b>x")
        self.assertContains(response, f"--scope-bg: {DEFAULT_COLOR}")

    def test_atencion_comes_first_and_the_menu_follows_the_permissions(self):
        top = self.page("matias.cavieres").context["menu_top"]
        self.assertEqual([item.label for item in top], ["Atención", "Dashboard"])
        top = self.page("coordinacion1").context["menu_top"]  # Coordinación no atiende público
        self.assertEqual([item.label for item in top], ["Dashboard"])

    def test_the_top_bar_has_the_user_name_a_link_to_the_profile_and_logout(self):
        response = self.page("matias.cavieres")
        self.assertEqual(response.context["me_name"], "Matías Cavieres")
        self.assertContains(response, 'class="topbar-label">Matías Cavieres</span>')
        self.assertContains(response, f'href="{reverse("core:profile")}"')
        self.assertContains(response, f'action="{reverse("core:logout")}"')
        self.assertContains(response, "Cerrar sesión")
        # Si la cuenta no tiene nombre, se muestra el usuario.
        self.assertEqual(self.page("admin").context["me_name"], "admin")

    def test_the_menu_marks_the_current_screen(self):
        response = self.page("matias.cavieres", "core:atencion")
        current = [item.label for item in response.context["menu_top"] if item.active]
        self.assertEqual(current, ["Atención"])

    def test_the_administration_group_lists_only_what_each_user_may_open(self):
        def administration(username):
            groups = self.page(username).context["menu_groups"]
            group = next(group for group in groups if group.key == "administracion")
            return {item.label for item in group.items}

        self.assertTrue({"Delegaciones", "Funcionarios", "Control de usuarios", "Auditoría"} <= administration("admin"))
        coordination = administration("coordinacion1")
        self.assertTrue({"Períodos", "Metas", "Ajustes"} <= coordination)
        self.assertFalse({"Delegaciones", "Control de usuarios", "Auditoría"} & coordination)
        officer = administration("matias.cavieres")
        self.assertFalse({"Delegaciones", "Control de usuarios", "Auditoría", "Períodos"} & officer)

    def test_the_menu_does_not_show_the_official_logo_or_claim_to_be_the_municipality_site(self):
        response = self.page("admin")
        self.assertContains(response, "no es un sitio oficial")
        self.assertContains(response, "Proyecto académico INACAP")
        self.assertNotContains(response, "laserena.cl")


class AtencionTests(PortalTestCase):
    def search(self, username, query="", **params):
        self.login_as(username)
        response = self.client.get(reverse("core:atencion"), {"q": query, **params})
        self.assertEqual(response.status_code, 200)
        return response

    def names(self, response):
        return [row.obj.contact_name for row in response.context["rows"]]

    def test_the_screen_has_the_search_box_the_buttons_and_the_latest_attentions(self):
        response = self.search("matias.cavieres")
        self.assertTemplateUsed(response, "core/atencion.html")
        for text in ("RUT / nombre / otro", "Buscar", "Exportar a Excel", "Registrar atención", "Últimas atenciones"):
            self.assertContains(response, text)
        self.assertEqual(set(self.names(response)), {"María Pérez", "Carla Muñoz"})

    def test_a_blank_search_shows_the_latest_attentions_too(self):
        response = self.search("matias.cavieres", "   ")
        self.assertContains(response, "Últimas atenciones")
        self.assertEqual(len(self.names(response)), 2)

    def test_search_by_rut_name_or_phone(self):
        for query in ("11.111.111-1", "11111111-1", "perez", "PÉREZ", "+56911111111"):
            with self.subTest(query=query):
                self.assertEqual(self.names(self.search("matias.cavieres", query)), ["María Pérez"])

    def test_a_wrong_check_digit_warns_but_still_searches(self):
        response = self.search("matias.cavieres", "11.111.111-2")
        self.assertTrue(response.context["wrong_rut"])
        self.assertContains(response, "dígito verificador")
        self.assertContains(response, "No encontramos atenciones")
        self.assertFalse(self.search("matias.cavieres", "11.111.111-1").context["wrong_rut"])
        self.assertFalse(self.search("matias.cavieres", "perez").context["wrong_rut"])

    def test_a_user_only_finds_attentions_of_its_own_delegation(self):
        self.assertEqual(self.names(self.search("matias.cavieres", "Ana Rojas")), [])  # Sur
        self.assertEqual(self.names(self.search("matias.quiroz", "Ana Rojas")), ["Ana Rojas"])

    def test_only_global_users_see_the_delegation_column(self):
        columns = [c.label for c in self.search("matias.cavieres").context["columns"]]
        self.assertEqual(
            columns,
            ["Identificador", "Fecha", "Contacto", "RUT", "Tipo de atención", "Sub atención", "Funcionario autor", "Estado"],
        )
        columns = [c.label for c in self.search("admin").context["columns"]]
        self.assertIn("Delegación", columns)
        self.assertEqual(columns[-1], "Estado")  # la delegación va antes del estado
        self.assertEqual(len(self.names(self.search("admin"))), 5)

    def test_every_row_can_start_a_new_attention_for_the_same_person(self):
        response = self.search("matias.cavieres", "11.111.111-1")
        beneficiary = self.activity("María Pérez").beneficiary
        link = f'{reverse("core:actividades_create")}?beneficiary={beneficiary.pk}'
        self.assertContains(response, link)
        form = self.client.get(link).context["form"]
        self.assertEqual(form["requester"].value(), beneficiary.rut)
        self.assertEqual(form["phone"].value(), beneficiary.phone)

    def test_the_new_attention_button_is_hidden_from_read_only_profiles(self):
        response = self.search("jefatura1")  # Jefatura consulta pero no registra
        self.assertNotContains(response, "Registrar atención")
        self.assertNotContains(response, "Nueva atención para")
        self.assertTrue(self.names(response))

    def test_without_permission_the_screen_is_refused(self):
        self.login_as("coordinacion1")
        self.assertEqual(self.client.get(reverse("core:atencion")).status_code, 403)
        self.client.logout()
        response = self.client.get(reverse("core:atencion"))
        self.assertEqual(response.status_code, 302)
        self.assertTrue(response["Location"].startswith(reverse("core:login") + "?next="))

    def test_the_excel_button_downloads_what_was_found(self):
        self.login_as("matias.cavieres")
        response = self.client.get(reverse("core:atencion"), {"q": "perez", "formato": "xlsx"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], export.CONTENT_TYPE)
        self.assertIn('filename="atenciones-', response["Content-Disposition"])
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet["A1"].value, "Atención")
        self.assertIn("Búsqueda: «perez»", sheet["A3"].value)
        headers = [cell.value for cell in sheet[export.HEADER_ROW]]
        row = [cell.value for cell in sheet[export.HEADER_ROW + 1]]
        self.assertEqual(row[headers.index("Contacto")], "María Pérez")
        self.assertEqual(row[headers.index("RUT")], "11111111-1")
        self.assertEqual(row[headers.index("Tipo de atención")], "Atención Social")
        self.assertEqual(row[headers.index("Sub atención")], "Visita Domiciliaria")
        self.assertEqual(row[headers.index("Identificador")], f"ATN{self.activity('María Pérez').pk:05d}")
        self.assertIsNone(sheet.cell(row=export.HEADER_ROW + 2, column=1).value)  # una sola fila
        log = AuditLog.objects.get(event="Listado exportado a Excel")
        self.assertEqual(log.affected_entity, "Activity")

    def test_the_search_box_is_checked_before_sending(self):
        response = self.search("matias.cavieres")
        self.assertContains(response, "data-validate")
        self.assertContains(response, 'id="atencion-q"')
        self.assertNotContains(response, "required data-label")  # ya no es obligatorio: se puede filtrar sin escribir
        self.assertContains(response, "data-skip-validate")  # pero el botón de Excel no lo exige


class ProfilePageTests(PortalTestCase):
    def test_profile_shows_account_delegation_officer_and_profiles(self):
        self.login_as("matias.cavieres")
        response = self.client.get(reverse("core:profile"))
        self.assertEqual(response.status_code, 200)
        for text in (
            "Mi perfil",
            "matias.cavieres",
            "Matías Cavieres",
            "Delegación Norte",
            "F-1001",
            "Encargado Social",
            "Funcionario",
            "Verificador",
            "Cambiar contraseña",
        ):
            self.assertContains(response, text)

    def test_global_users_without_officer_see_a_notice(self):
        self.login_as("coordinacion1")
        response = self.client.get(reverse("core:profile"))
        self.assertContains(response, "Todas las delegaciones")
        self.assertContains(response, "no tiene un funcionario asociado")

    def test_anonymous_visitors_are_sent_to_the_login(self):
        self.assertEqual(self.client.get(reverse("core:profile")).status_code, 302)

    def test_changing_the_password_keeps_the_session_and_is_audited(self):
        user = self.login_as("matias.cavieres")
        response = self.client.post(
            reverse("core:profile"),
            {"old_password": PASSWORD, "new_password1": "OtraClave#2026x", "new_password2": "OtraClave#2026x"},
        )
        self.assertRedirects(response, reverse("core:profile"))
        user.refresh_from_db()
        self.assertTrue(user.check_password("OtraClave#2026x"))
        self.assertEqual(self.client.get(reverse("core:dashboard")).status_code, 200)  # sigue con la sesión abierta
        log = AuditLog.objects.get(event="Contraseña cambiada")
        self.assertEqual(log.user, user)
        self.assertEqual(log.record_id, user.pk)
        self.assertFalse(log.new_value)  # la contraseña no queda en ninguna parte

    def test_a_wrong_current_password_or_a_mismatch_changes_nothing(self):
        user = self.login_as("matias.cavieres")
        response = self.client.post(
            reverse("core:profile"),
            {"old_password": "equivocada", "new_password1": "OtraClave#2026x", "new_password2": "OtraClave#2026x"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("old_password", response.context["form"].errors)
        response = self.client.post(
            reverse("core:profile"),
            {"old_password": PASSWORD, "new_password1": "OtraClave#2026x", "new_password2": "no coincide"},
        )
        self.assertIn("new_password2", response.context["form"].errors)
        user.refresh_from_db()
        self.assertTrue(user.check_password(PASSWORD))
        self.assertFalse(AuditLog.objects.filter(event="Contraseña cambiada").exists())

    def test_a_weak_new_password_is_refused(self):
        user = self.login_as("matias.cavieres")
        response = self.client.post(
            reverse("core:profile"),
            {"old_password": PASSWORD, "new_password1": "12345678", "new_password2": "12345678"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("new_password1", response.context["form"].errors)
        self.assertTrue(User.objects.get(pk=user.pk).check_password(PASSWORD))

    def test_the_form_says_contrasena(self):
        self.login_as("matias.cavieres")
        response = self.client.get(reverse("core:profile"))
        labels = [str(field.label) for field in response.context["form"]]
        self.assertEqual(labels, ["Contraseña actual", "Contraseña nueva", "Repite la contraseña nueva"])


class AdministrationHubTests(PortalTestCase):
    def test_the_hub_lists_what_the_user_can_administer_without_linking_itself(self):
        self.login_as("admin")
        response = self.client.get(reverse("core:administration"))
        self.assertEqual(response.status_code, 200)
        for text in ("Delegaciones", "Funcionarios", "Control de usuarios", "Auditoría"):
            self.assertContains(response, text)
        titles = [title for title, _ in response.context["sections"]]
        self.assertEqual(titles, ["Organización", "Planificación", "Catálogos", "Sistema"])
        links = [item.url for _, items in response.context["sections"] for item in items]
        self.assertNotIn(reverse("core:administration"), links)

    def test_the_hub_is_the_first_entry_of_the_administration_group(self):
        self.login_as("admin")
        response = self.client.get(reverse("core:dashboard"))
        group = next(group for group in response.context["menu_groups"] if group.key == "administracion")
        self.assertEqual(group.sections[0][1][0].label, "Resumen")

    def test_users_with_nothing_to_administer_are_refused(self):
        self.make_user("consulta.norte", groups=["Consulta"], officer=self.officer("F-1004"))
        self.assertTrue(self.client.login(username="consulta.norte", password=PASSWORD))
        self.assertEqual(self.client.get(reverse("core:administration")).status_code, 403)

    def test_coordination_sees_only_its_own_maintainers(self):
        self.login_as("coordinacion1")
        response = self.client.get(reverse("core:administration"))
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, "Control de usuarios")
