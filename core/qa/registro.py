"""Arma el registro de pruebas funcionales (A3) y lo escribe como Excel y CSV."""

import csv
import json
from collections import Counter, OrderedDict
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.qa import defectos, pendientes
from core.qa.catalogo import CLASES, ESPERADO, REQUISITOS
from core.qa.pasos import Indice

ZONA = ZoneInfo("America/Santiago")
APROBADA, FALLIDA, ERROR, OMITIDA, SIN_EJECUTAR = "Aprobada", "Fallida", "Error", "Omitida", "No ejecutada"

COLUMNAS = [
    ("ID", 10),
    ("Requisito", 16),
    ("Módulo", 34),
    ("Prueba (clase.método)", 52),
    ("Precondiciones", 46),
    ("Pasos", 58),
    ("Resultado esperado", 60),
    ("Resultado obtenido", 52),
    ("Estado", 13),
    ("Evidencia", 52),
    ("Responsable", 24),
    ("Fecha", 18),
]


def identificador(modulo, clase, metodo):
    """Nombre completo de la prueba, el mismo que usa Django: ``core.test_a1.Clase.test_metodo``."""
    return f"{modulo}.{clase}.test_{metodo}"


def pruebas_del_proyecto(base):
    """Todas las pruebas automáticas: lista de (archivo, clase, método) en el orden de los archivos."""
    indice = Indice(base)
    pruebas = []
    for clase, metodos in ESPERADO.items():
        archivo = indice.archivo(clase)
        pruebas.extend((archivo, clase, metodo) for metodo in metodos)
    pruebas.sort(key=lambda p: (p[0], indice.clases[p[1]][1].lineno, list(ESPERADO[p[1]]).index(p[2])))
    return indice, pruebas


def filas(base, resultados, responsable, fecha, archivo_ejecucion):
    """Una fila por prueba. ``resultados`` mapea el id de Django a ``{"estado": ..., "detalle": ...}``."""
    indice, pruebas = pruebas_del_proyecto(base)
    salida = []
    for numero, (archivo, clase, metodo) in enumerate(pruebas, start=1):
        modulo_py = archivo[:-3].replace("/", ".")
        resultado = resultados.get(identificador(modulo_py, clase, metodo), {"estado": SIN_EJECUTAR, "detalle": ""})
        requisitos, modulo, precondicion = CLASES[clase]
        estado = resultado["estado"]
        if estado == APROBADA:
            obtenido = "Se obtuvo lo esperado: la prueba automática terminó sin errores."
        elif estado == OMITIDA:
            obtenido = f"No se ejecutó (omitida). Motivo: {resultado.get('detalle') or 'no indicado'}"
        elif estado == SIN_EJECUTAR:
            obtenido = "Sin resultado: la prueba no se ejecutó en esta corrida."
        else:
            obtenido = "No se obtuvo lo esperado. " + (resultado.get("detalle") or "").strip().splitlines()[-1][:300]
        salida.append(
            OrderedDict(
                [
                    ("ID", f"RP-{numero:04d}"),
                    ("Requisito", ", ".join(requisitos)),
                    ("Módulo", modulo),
                    ("Prueba (clase.método)", f"{clase}.test_{metodo}"),
                    ("Precondiciones", precondicion),
                    ("Pasos", "\n".join(f"{n}. {p}" for n, p in enumerate(indice.pasos(clase, metodo), start=1))),
                    ("Resultado esperado", ESPERADO[clase][metodo]),
                    ("Resultado obtenido", obtenido),
                    ("Estado", estado),
                    ("Evidencia", f"{archivo}::{clase}.test_{metodo} · {archivo_ejecucion}"),
                    ("Responsable", responsable),
                    ("Fecha", fecha),
                ]
            )
        )
    return salida


def resumen_por_requisito(registro):
    cuentas = {codigo: Counter() for codigo in REQUISITOS}
    for fila in registro:
        for codigo in fila["Requisito"].split(", "):
            cuentas[codigo][fila["Estado"]] += 1
    return cuentas


def _estilo_encabezado(hoja, fila=1):
    for celda in hoja[fila]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="920000")
        celda.alignment = Alignment(wrap_text=True, vertical="center")


def _hoja_tabla(libro, titulo, encabezados, anchos, datos, wrap=True):
    hoja = libro.create_sheet(titulo)
    hoja.append(encabezados)
    for fila in datos:
        hoja.append(list(fila))
    for i, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(i)].width = ancho
    _estilo_encabezado(hoja)
    hoja.freeze_panes = "A2"
    if datos:
        hoja.auto_filter.ref = f"A1:{get_column_letter(len(encabezados))}{len(datos) + 1}"
    if wrap:
        for fila in hoja.iter_rows(min_row=2):
            for celda in fila:
                celda.alignment = Alignment(wrap_text=True, vertical="top")
    return hoja


def escribir(salida, registro, meta, extras):
    """Escribe ``registro_pruebas_funcionales.xlsx`` y ``.csv`` en ``salida`` (carpeta)."""
    salida = Path(salida)
    salida.mkdir(parents=True, exist_ok=True)
    cuentas = resumen_por_requisito(registro)
    total = Counter(f["Estado"] for f in registro)

    libro = Workbook()
    resumen = libro.active
    resumen.title = "Resumen"
    resumen.append(["SGR · Registro de pruebas funcionales (A3)"])
    resumen["A1"].font = Font(bold=True, size=14)
    for clave, valor in meta.items():
        resumen.append([clave, valor])
    resumen.append([])
    resumen.append(["Estado", "Pruebas"])
    _estilo_encabezado(resumen, resumen.max_row)
    for estado in (APROBADA, FALLIDA, ERROR, OMITIDA, SIN_EJECUTAR):
        resumen.append([estado, total.get(estado, 0)])
    resumen.append(["Total", len(registro)])
    resumen.append([])
    resumen.append(
        [
            "Cómo leerlo: cada fila de «Registro» es una prueba automática real del proyecto (Django). Los pasos se extraen "
            "del código de la prueba; el detalle completo está en el archivo indicado en «Evidencia». Las pruebas con "
            "navegador (hoja «Navegador») y las de estrés (hoja «Estrés») se registran aparte. Lo que no está probado "
            "figura en «Pendientes» con su justificación: no se declara aprobado."
        ]
    )
    resumen.column_dimensions["A"].width = 38
    resumen.column_dimensions["B"].width = 90
    resumen[f"A{resumen.max_row}"].alignment = Alignment(wrap_text=True, vertical="top")
    resumen.row_dimensions[resumen.max_row].height = 90

    hoja = _hoja_tabla(libro, "Registro", [c for c, _ in COLUMNAS], [a for _, a in COLUMNAS], [list(f.values()) for f in registro])
    verde, rojo = PatternFill("solid", fgColor="D9EAD3"), PatternFill("solid", fgColor="F4CCCC")
    for fila in hoja.iter_rows(min_row=2):
        fila[8].fill = verde if fila[8].value == APROBADA else rojo

    datos = []
    for codigo, (seccion, texto) in REQUISITOS.items():
        c = cuentas[codigo]
        n = sum(c.values())
        if n == 0:
            estado = "Sin prueba automática (ver «Pendientes»)"
        elif c[APROBADA] == n:
            estado = "Cubierto: todas aprobadas"
        else:
            estado = "Con pruebas que no pasan"
        datos.append([codigo, seccion, texto, n, c[APROBADA], n - c[APROBADA], estado])
    _hoja_tabla(libro, "Requisitos", ["Código", "Sección", "Requisito", "Pruebas", "Aprobadas", "No aprobadas", "Estado"], [10, 12, 80, 10, 11, 13, 40], datos)

    _hoja_tabla(
        libro,
        "Navegador",
        ["ID", "Requisito", "Prueba", "Pasos", "Resultado esperado", "Resultado obtenido", "Estado", "Evidencia", "Responsable", "Fecha"],
        [10, 14, 44, 60, 56, 52, 12, 46, 24, 18],
        [[f.get(k, "") for k in ("id", "requisito", "prueba", "pasos", "esperado", "obtenido", "estado", "evidencia", "responsable", "fecha")] for f in extras.get("navegador", [])]
        or [["—", "", "Todavía no se ejecutaron pruebas con navegador. Ver docs/evidencias/README.md.", "", "", "", SIN_EJECUTAR, "", "", ""]],
    )
    _hoja_tabla(
        libro,
        "Estrés",
        ["Escenario", "Concurrencia", "Duración (s)", "Solicitudes", "Req/s", "Latencia prom. (ms)", "p95 (ms)", "Errores", "Entorno", "Veredicto"],
        [34, 13, 13, 12, 10, 18, 11, 10, 46, 40],
        extras.get("estres", []) or [["Sin resultados cargados", "", "", "", "", "", "", "", "", "Pendiente: se ejecuta contra AWS (docs/pruebas/estres/)"]],
    )
    _hoja_tabla(libro, "Defectos", defectos.COLUMNAS, defectos.ANCHOS, defectos.filas())
    _hoja_tabla(libro, "Pendientes", pendientes.COLUMNAS, pendientes.ANCHOS, pendientes.filas())
    libro.save(salida / "registro_pruebas_funcionales.xlsx")

    with open(salida / "registro_pruebas_funcionales.csv", "w", newline="", encoding="utf-8-sig") as archivo:
        escritor = csv.writer(archivo)
        escritor.writerow([c for c, _ in COLUMNAS])
        for f in registro:
            escritor.writerow(list(f.values()))
    return total


def cargar_json(ruta):
    ruta = Path(ruta)
    return json.loads(ruta.read_text(encoding="utf-8")) if ruta.exists() else None


def ahora():
    return datetime.now(ZONA).strftime("%Y-%m-%d %H:%M")
