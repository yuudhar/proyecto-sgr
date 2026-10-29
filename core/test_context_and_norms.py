import struct
from io import BytesIO
from pathlib import Path
from unittest import mock

from django.contrib.auth.models import Group, User
from django.contrib.staticfiles import finders
from django.urls import reverse
from openpyxl import load_workbook

from accounts.delegations import extra_delegations, listed_extra_delegations, sync_extra_delegations
from accounts.models import RoleProfile, UserRoleDelegation
from core.access import SESSION_DELEGATION
from core.models import AuditLog
from core.testing import PASSWORD, PortalTestCase
from operations.models import Activity, Commitment
from organization.models import Delegation, Officer, Territory


class ContextTestCase(PortalTestCase):

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.center = Delegation.objects.create(name="Delegación Centro", color="#6b4e9b")

    def switch(self, value, next_url=""):
        return self.client.post(reverse("core:set_context"), {"delegacion": value, "next": next_url})

    def header(self, url=None):
        response = self.client.get(url or reverse("core:dashboard"))
        html = response.content.decode()
        start = html.index('<header class="topbar">')
        return html[start : html.index("</header>", start)]

    def select_options(self, header):
        start = header.index('<select')
        end = header.index("</select>", start)
        import re

        return re.findall(r"<option value=\"([^\"]*)\"[^>]*>([^<]*)</option>", header[start:end])

    def contacts(self, response):
        return {row.obj.contact_name for row in response.context["rows"]}

    def contacts_of(self, delegation):
        return set(
            Activity.objects.filter(delegation=delegation, deleted_at__isnull=True).values_list("contact_name", flat=True)
        )

    def two_delegation_worker(self, username="func.dos"):
        user = self.make_user(username, ["Funcionario"], self.officer("F-1004"))
        sync_extra_delegations(user, [self.south])
        return user


class AuthorizedDelegationsTests(ContextTestCase):

    def test_a_user_with_one_delegation_has_a_fixed_context_without_selector(self):
        self.login_as("matias.cavieres")
        header = self.header()
        self.assertNotIn("<select", header)
        self.assertNotIn("data-context-form", header)
        self.assertIn(f'<span class="topbar-code">DEL{self.north.pk:02d}</span>', header)
        scope = self.client.get(reverse("core:dashboard")).context["me_scope"]
        self.assertFalse(scope["choose"])
        self.assertEqual(scope["options"][0]["label"], f"DEL{self.north.pk:02d} · Delegación Norte")

    def test_the_selector_only_offers_the_authorized_delegations(self):
        self.login_as("supervision1")
        header = self.header()
        options = self.select_options(header)
        self.assertEqual(
            options,
            [
                (str(self.north.pk), f"DEL{self.north.pk:02d} · Delegación Norte"),
                (str(self.south.pk), f"DEL{self.south.pk:02d} · Delegación Sur"),
            ],
        )
        self.assertNotIn("Delegación Centro", header)
        self.assertNotIn("Todas las delegaciones", header)
        self.assertIn(f'<option value="{self.north.pk}" selected>', header)

    def test_a_user_who_sees_everything_gets_all_the_delegations_and_starts_in_all(self):
        self.login_as("coordinacion1")
        header = self.header()
        options = self.select_options(header)
        self.assertEqual(
            [value for value, _label in options],
            ["todas", str(self.north.pk), str(self.south.pk), str(self.center.pk)],
        )
        self.assertIn('<option value="todas" selected>', header)
        self.assertNotIn("topbar-code", header)

    def test_archived_and_inactive_delegations_are_not_offered(self):
        Delegation.objects.filter(pk=self.center.pk).update(deleted_at="2026-01-01T00:00:00Z")
        self.login_as("admin")
        values = [value for value, _label in self.select_options(self.header())]
        self.assertEqual(values, ["todas", str(self.north.pk), str(self.south.pk)])

        Delegation.objects.filter(pk=self.south.pk).update(status=Delegation.Status.INACTIVE)
        self.login_as("supervision1")
        self.assertNotIn("<select", self.header())

    def test_extra_rows_only_count_while_the_user_keeps_an_active_role(self):
        user = User.objects.get(username="supervision1")
        self.assertEqual([d.pk for d in extra_delegations(user)], [self.south.pk])

        RoleProfile.objects.filter(group__name="Jefatura").update(status=RoleProfile.Status.INACTIVE)
        self.assertEqual(list(extra_delegations(user)), [])
        RoleProfile.objects.filter(group__name="Jefatura").update(status=RoleProfile.Status.ACTIVE)

        user.groups.remove(Group.objects.get(name="Jefatura"))
        self.assertEqual(list(extra_delegations(user)), [])
        self.assertTrue(UserRoleDelegation.objects.filter(user=user).exists())

    def test_the_list_version_agrees_with_the_query_version(self):
        user = User.objects.get(username="supervision1")
        group = Group.objects.get(name="Jefatura")
        states = [
            lambda: None,
            lambda: RoleProfile.objects.filter(group=group).update(status=RoleProfile.Status.INACTIVE),
            lambda: RoleProfile.objects.filter(group=group).update(status=RoleProfile.Status.ACTIVE),
            lambda: Delegation.objects.filter(pk=self.south.pk).update(status=Delegation.Status.INACTIVE),
            lambda: Delegation.objects.filter(pk=self.south.pk).update(status=Delegation.Status.ACTIVE),
            lambda: Delegation.objects.filter(pk=self.south.pk).update(deleted_at="2026-01-01T00:00:00Z"),
            lambda: Delegation.objects.filter(pk=self.south.pk).update(deleted_at=None),
            lambda: UserRoleDelegation.objects.filter(user=user).update(deleted_at="2026-01-01T00:00:00Z"),
            lambda: UserRoleDelegation.objects.filter(user=user).update(deleted_at=None),
            lambda: user.groups.remove(group),
        ]
        for change in states:
            change()
            fresh = User.objects.prefetch_related("groups__role_profile", "role_delegations__delegation").get(pk=user.pk)
            self.assertEqual(
                [d.pk for d in listed_extra_delegations(fresh)],
                [d.pk for d in extra_delegations(user)],
            )

    def test_a_row_for_a_role_the_user_does_not_have_is_rejected_by_the_model(self):
        from django.core.exceptions import ValidationError

        user = User.objects.get(username="matias.cavieres")
        row = UserRoleDelegation(user=user, role=Group.objects.get(name="Jefatura"), delegation=self.south)
        with self.assertRaises(ValidationError):
            row.full_clean()

    def test_the_combination_is_unique(self):
        from django.db import IntegrityError, transaction

        user = User.objects.get(username="supervision1")
        with self.assertRaises(IntegrityError), transaction.atomic():
            UserRoleDelegation.objects.create(user=user, role=Group.objects.get(name="Jefatura"), delegation=self.south)

    def test_syncing_adds_removes_and_revives(self):
        user = self.make_user("sync.uno", ["Funcionario", "Verificador"], self.officer("F-1004"))
        self.assertEqual(sync_extra_delegations(user, [self.south, self.center]), (4, 0))
        self.assertEqual(sync_extra_delegations(user, [self.south, self.center]), (0, 0))
        self.assertEqual(sync_extra_delegations(user, [self.south]), (0, 2))
        UserRoleDelegation.objects.filter(user=user).update(deleted_at="2026-01-01T00:00:00Z")
        self.assertEqual(sync_extra_delegations(user, [self.south]), (2, 0))
        self.assertEqual(UserRoleDelegation.objects.filter(user=user, deleted_at__isnull=True).count(), 2)


class SwitchContextTests(ContextTestCase):

    def test_switching_saves_the_choice_in_the_session_and_audits_it(self):
        user = self.login_as("supervision1")
        response = self.switch(self.south.pk, reverse("core:atencion"))
        self.assertRedirects(response, reverse("core:atencion"))
        self.assertEqual(self.client.session[SESSION_DELEGATION], self.south.pk)

        log = AuditLog.objects.get(event="Cambio de delegación activa")
        self.assertEqual(log.user, user)
        self.assertEqual(log.previous_value, f"DEL{self.north.pk:02d} · Delegación Norte")
        self.assertEqual(log.new_value, f"DEL{self.south.pk:02d} · Delegación Sur")

        header = self.header()
        self.assertIn(f'<option value="{self.south.pk}" selected>', header)
        scope = self.client.get(reverse("core:dashboard")).context["me_scope"]
        self.assertEqual(scope["code"], f"DEL{self.south.pk:02d}")
        self.assertEqual(scope["name"], "Delegación Sur")

    def test_the_new_delegation_is_announced_with_a_toast(self):
        self.login_as("supervision1")
        response = self.switch(self.south.pk, reverse("core:dashboard"), )
        page = self.client.get(response["Location"])
        messages = [(m.message, m.extra_tags) for m in page.context["messages"]]
        self.assertEqual(len(messages), 1)
        self.assertIn("Ahora trabajas en: DEL", messages[0][0])
        self.assertIn("Delegación Sur", messages[0][0])
        self.assertEqual(messages[0][1], "toast")

    def test_choosing_the_same_delegation_changes_nothing(self):
        self.login_as("supervision1")
        self.switch(self.north.pk)
        self.assertFalse(AuditLog.objects.filter(event="Cambio de delegación activa").exists())

    def test_a_forged_id_is_a_403_and_is_audited(self):
        user = self.login_as("supervision1")
        for forged in (self.center.pk, "todas", "9999", "abc", "-1", "1; DROP TABLE", ""):
            with self.subTest(forged=forged):
                response = self.switch(forged)
                self.assertEqual(response.status_code, 403)
                self.assertNotIn(SESSION_DELEGATION, self.client.session)
        self.assertEqual(AuditLog.objects.filter(event="Cambio de delegación rechazado", user=user).count(), 7)
        rejected = AuditLog.objects.filter(event="Cambio de delegación rechazado").first()
        self.assertTrue(rejected.new_value.startswith("Delegación pedida: "))

    def test_a_user_with_one_delegation_cannot_switch_even_by_hand(self):
        self.login_as("matias.cavieres")
        for forged in (self.south.pk, self.center.pk, "todas"):
            with self.subTest(forged=forged):
                self.assertEqual(self.switch(forged).status_code, 403)
                self.assertNotIn(SESSION_DELEGATION, self.client.session)
        self.assertEqual(self.switch(self.north.pk).status_code, 302)

    def test_a_403_page_still_shows_the_header(self):
        self.login_as("matias.cavieres")
        response = self.switch(self.south.pk)
        self.assertContains(response, "No tienes autorizada esa delegación.", status_code=403)
        self.assertContains(response, 'class="topbar-code"', status_code=403)

    def test_only_post_is_accepted_and_login_is_required(self):
        self.login_as("supervision1")
        self.assertEqual(self.client.get(reverse("core:set_context")).status_code, 405)
        self.client.logout()
        response = self.switch(self.south.pk)
        self.assertEqual(response.status_code, 302)
        self.assertTrue(response["Location"].startswith(reverse("core:login")))

    def test_the_csrf_token_is_required(self):
        from django.test import Client

        user = User.objects.get(username="supervision1")
        strict = Client(enforce_csrf_checks=True)
        strict.force_login(user)
        response = strict.post(reverse("core:set_context"), {"delegacion": self.south.pk})
        self.assertEqual(response.status_code, 403)
        self.assertNotIn(SESSION_DELEGATION, strict.session)

    def test_a_global_user_can_pick_one_delegation_and_go_back_to_all(self):
        self.login_as("coordinacion1")
        self.switch(self.south.pk)
        self.assertEqual(self.client.session[SESSION_DELEGATION], self.south.pk)
        scope = self.client.get(reverse("core:dashboard")).context["me_scope"]
        self.assertFalse(scope["viewing_all"])
        self.assertTrue(scope["all"])
        self.assertEqual(scope["note"], "Con acceso a todas las delegaciones")

        self.switch("todas")
        self.assertEqual(self.client.session[SESSION_DELEGATION], "todas")
        scope = self.client.get(reverse("core:dashboard")).context["me_scope"]
        self.assertTrue(scope["viewing_all"])
        self.assertEqual(scope["name"], "Todas las delegaciones")

    def test_the_choice_lasts_for_the_session_and_each_login_starts_in_the_own_delegation(self):
        self.login_as("supervision1")
        self.switch(self.south.pk)
        for name in ("core:dashboard", "core:atencion", "core:profile"):
            self.assertEqual(self.client.get(reverse(name)).context["me_scope"]["name"], "Delegación Sur")
        self.client.logout()
        self.assertNotIn(SESSION_DELEGATION, self.client.session)
        response = self.client.post(reverse("core:login"), {"username": "supervision1", "password": PASSWORD})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(self.client.get(reverse("core:dashboard")).context["me_scope"]["name"], "Delegación Norte")

    def test_logging_in_again_over_the_same_session_forgets_the_choice(self):
        from django.contrib.auth import login
        from django.contrib.sessions.backends.db import SessionStore
        from django.test import RequestFactory

        request = RequestFactory().get("/")
        request.session = SessionStore()
        request.session[SESSION_DELEGATION] = self.south.pk
        user = User.objects.get(username="supervision1")
        login(request, user, backend="django.contrib.auth.backends.ModelBackend")
        self.assertNotIn(SESSION_DELEGATION, request.session)

    def test_a_delegation_that_stops_being_authorized_is_forgotten_on_the_next_request(self):
        user = self.login_as("supervision1")
        self.switch(self.south.pk)
        self.assertEqual(self.client.get(reverse("core:dashboard")).context["me_scope"]["name"], "Delegación Sur")
        sync_extra_delegations(user, [])
        response = self.client.get(reverse("core:atencion"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.contacts(response), self.contacts_of(self.north))
        self.assertNotIn(SESSION_DELEGATION, self.client.session)

    def test_a_tampered_session_value_is_ignored(self):
        self.login_as("matias.cavieres")
        session = self.client.session
        session[SESSION_DELEGATION] = self.south.pk
        session.save()
        response = self.client.get(reverse("core:atencion"))
        self.assertEqual(self.contacts(response), self.contacts_of(self.north))
        self.assertNotIn(SESSION_DELEGATION, self.client.session)
        for junk in ("todas", True, [1], {"a": 1}, None, 10**30):
            session = self.client.session
            session[SESSION_DELEGATION] = junk
            session.save()
            self.assertEqual(self.client.get(reverse("core:dashboard")).status_code, 200)

    def test_where_to_go_after_switching(self):
        self.login_as("supervision1")
        list_url = reverse("core:atencion")
        activity = Activity.objects.filter(delegation=self.north, deleted_at__isnull=True).first()
        cases = {
            f"{list_url}?q=perez&page=3": list_url,
            reverse("core:actividades_create"): reverse("core:actividades_create"),
            reverse("core:actividades_detail", args=[activity.pk]): list_url,
            reverse("core:actividades_update", args=[activity.pk]): list_url,
            "https://example.com/robo": reverse("core:dashboard"),
            "//example.com/robo": reverse("core:dashboard"),
            "/no-existe-esta-pagina/": reverse("core:dashboard"),
            reverse("core:set_context"): reverse("core:dashboard"),
            "": reverse("core:dashboard"),
            "javascript:alert(1)": reverse("core:dashboard"),
        }
        for next_url, expected in cases.items():
            with self.subTest(next_url=next_url):
                self.switch(self.north.pk)
                response = self.switch(self.south.pk, next_url)
                self.assertEqual(response.status_code, 302)
                self.assertEqual(response["Location"], expected)
                self.assertEqual(self.client.session[SESSION_DELEGATION], self.south.pk)
                self.switch(self.north.pk)

    def test_after_validating_an_evidence_screen_the_switch_goes_to_the_list(self):
        user = self.make_user("revisa.dos", ["Funcionario", "Verificador"], self.officer("F-1004"))
        sync_extra_delegations(user, [self.south])
        self.client.force_login(user)
        from evidence.models import Evidence

        evidence = Evidence.objects.filter(activity__delegation=self.north, deleted_at__isnull=True).first()
        response = self.switch(self.south.pk, reverse("core:validate_evidence", args=[evidence.pk]))
        self.assertEqual(response["Location"], reverse("core:validate_list"))


class ContextFollowsEverywhereTests(ContextTestCase):

    def test_the_lists_follow_the_active_delegation(self):
        self.login_as("supervision1")
        url = reverse("core:atencion")
        self.assertEqual(self.contacts(self.client.get(url)), self.contacts_of(self.north))
        self.switch(self.south.pk)
        self.assertEqual(self.contacts(self.client.get(url)), self.contacts_of(self.south))
        self.switch(self.north.pk)
        self.assertEqual(self.contacts(self.client.get(url)), self.contacts_of(self.north))

    def test_the_other_delegation_records_are_not_found_by_address(self):
        self.login_as("supervision1")
        south_activity = Activity.objects.filter(delegation=self.south, deleted_at__isnull=True).first()
        detail = reverse("core:actividades_detail", args=[south_activity.pk])
        self.assertEqual(self.client.get(detail).status_code, 404)
        self.switch(self.south.pk)
        self.assertEqual(self.client.get(detail).status_code, 200)

    def test_all_the_lists_that_belong_to_a_delegation_follow_it(self):
        self.login_as("admin")
        for name, lookup in (
            ("core:agenda_list", "delegation"),
            ("core:evidencias_list", "activity__delegation"),
            ("core:territorios_list", "delegation"),
            ("core:funcionarios_list", "delegation"),
        ):
            with self.subTest(name=name):
                self.switch("todas")
                everything = {row.obj.pk for row in self.client.get(reverse(name)).context["rows"]}
                self.switch(self.south.pk)
                south = {row.obj.pk for row in self.client.get(reverse(name)).context["rows"]}
                crud_model = self.client.get(reverse(name)).context["crud"].model
                expected = set(crud_model.objects.filter(**{lookup: self.south}, deleted_at__isnull=True).values_list("pk", flat=True))
                self.assertEqual(south, expected)
                self.assertLessEqual(south, everything)

    def test_the_delegation_filter_disappears_when_one_delegation_is_active(self):
        self.login_as("admin")
        url = reverse("core:atencion")
        names = [spec.name for spec in self.client.get(url).context["filters"]]
        self.assertIn("delegation", names)
        self.switch(self.south.pk)
        names = [spec.name for spec in self.client.get(url).context["filters"]]
        self.assertNotIn("delegation", names)
        response = self.client.get(url, {"delegation": self.north.pk})
        self.assertEqual(self.contacts(response), self.contacts_of(self.south))

    def test_the_attention_screen_shows_the_delegation_column_only_when_looking_at_all(self):
        self.login_as("admin")
        columns = [column.label for column in self.client.get(reverse("core:atencion")).context["columns"]]
        self.assertIn("Delegación", columns)
        self.switch(self.south.pk)
        columns = [column.label for column in self.client.get(reverse("core:atencion")).context["columns"]]
        self.assertNotIn("Delegación", columns)

    def test_the_excel_export_follows_the_delegation_and_says_which_one(self):
        self.login_as("supervision1")
        url = reverse("core:atencion")

        def sheet(**params):
            response = self.client.get(url, {"formato": "xlsx", **params})
            self.assertEqual(response.status_code, 200)
            return load_workbook(BytesIO(response.content)).active

        north = sheet()
        self.assertIn(f"DEL{self.north.pk:02d} · Delegación Norte", north["A3"].value)
        self.switch(self.south.pk)
        south = sheet()
        self.assertIn(f"Delegación: DEL{self.south.pk:02d} · Delegación Sur", south["A3"].value)
        contacts = {row[2] for row in south.iter_rows(min_row=6, values_only=True) if row[2]}
        self.assertEqual(contacts, self.contacts_of(self.south))
        self.assertTrue(contacts.isdisjoint(self.contacts_of(self.north)))

    def test_an_export_that_sees_everything_has_no_delegation_line(self):
        self.login_as("admin")
        response = self.client.get(reverse("core:atencion"), {"formato": "xlsx"})
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet["A3"].value, "Sin filtros")

    def test_a_global_user_exporting_one_delegation_gets_only_its_rows(self):
        self.login_as("admin")
        self.switch(self.south.pk)
        response = self.client.get(reverse("core:atencion"), {"formato": "xlsx"})
        sheet = load_workbook(BytesIO(response.content)).active
        contacts = {row[2] for row in sheet.iter_rows(min_row=6, values_only=True) if row[2]}
        self.assertEqual(contacts, self.contacts_of(self.south))
        log = AuditLog.objects.filter(event="Listado exportado a Excel").latest("pk")
        self.assertIn("Delegación: DEL", log.new_value)

    def test_the_dashboard_counts_follow_the_delegation(self):
        self.login_as("admin")

        def open_commitments():
            cards = self.client.get(reverse("core:dashboard")).context["pending"]
            return next(card["total"] for card in cards if card["label"] == "Compromisos abiertos")

        everything = open_commitments()
        self.switch(self.south.pk)
        south = open_commitments()
        self.switch(self.north.pk)
        north = open_commitments()
        self.assertEqual(south + north, everything)

    def test_the_board_and_reports_follow_the_delegation(self):
        self.login_as("supervision1")
        board = self.client.get(reverse("core:delegation_board"))
        self.assertEqual(board.context["delegation"], self.north)
        self.assertFalse(board.context["can_pick_delegation"])
        self.switch(self.south.pk)
        board = self.client.get(reverse("core:delegation_board"))
        self.assertEqual(board.context["delegation"], self.south)
        reports = self.client.get(reverse("core:reports"))
        self.assertEqual(reports.context["delegation"], self.south)
        self.assertFalse(reports.context["can_pick_delegation"])

    def test_a_global_user_looking_at_all_keeps_the_local_picker_of_the_board(self):
        self.login_as("admin")
        board = self.client.get(reverse("core:delegation_board"))
        self.assertTrue(board.context["can_pick_delegation"])
        self.switch(self.south.pk)
        self.assertFalse(self.client.get(reverse("core:delegation_board")).context["can_pick_delegation"])

    def test_the_indicators_follow_the_delegation(self):
        self.login_as("admin")
        url = reverse("core:indicadores_list")
        self.switch(self.south.pk)
        officers = {row.obj.goal.officer.delegation for row in self.client.get(url).context["rows"] if row.obj.goal.officer_id}
        self.assertEqual(officers, {self.south})

    def test_the_validation_list_follows_the_delegation(self):
        user = self.make_user("revisa.dos", ["Funcionario", "Verificador"], self.officer("F-1004"))
        sync_extra_delegations(user, [self.south])
        self.client.force_login(user)
        north_codes = {item["evidence"].unique_code for item in self.client.get(reverse("core:validate_list")).context["items"]}
        self.assertTrue(all(code.startswith("EVD-NORTE-") for code in north_codes))
        self.switch(self.south.pk)
        south_codes = {item["evidence"].unique_code for item in self.client.get(reverse("core:validate_list")).context["items"]}
        self.assertTrue(south_codes)
        self.assertTrue(all(code.startswith("EVD-SUR-") for code in south_codes))

    def test_the_dashboard_says_where_the_person_is_working(self):
        self.login_as("supervision1")
        self.assertContains(self.client.get(reverse("core:dashboard")), "Estás trabajando en <strong>Delegación Norte</strong>")
        self.switch(self.south.pk)
        self.assertContains(self.client.get(reverse("core:dashboard")), "Estás trabajando en <strong>Delegación Sur</strong>")
        self.login_as("matias.cavieres")
        self.assertNotContains(self.client.get(reverse("core:dashboard")), "Estás trabajando en")

    def test_the_profile_lists_the_authorized_delegations(self):
        self.login_as("supervision1")
        page = self.client.get(reverse("core:profile"))
        self.assertContains(page, "Delegación activa")
        self.assertContains(page, "Delegaciones autorizadas")
        self.assertContains(page, "Delegación Norte, DEL")


class ContextInFormsTests(ContextTestCase):

    def commitment_data(self, **changes):
        territory = Territory.objects.filter(delegation=self.south).first()
        data = {
            "requester": "Junta de vecinos",
            "requester_type": Commitment.RequesterType.values[0],
            "source": "",
            "request_date": "2026-10-01",
            "description": "Arreglo de luminarias.",
            "territory": territory.pk,
            "due_date": "2026-12-01",
            "support": "",
            "status": Commitment.Status.values[0],
            "notes": "",
            "responsible_officer": Officer.objects.filter(delegation=self.south).first().pk,
            "delegation": self.south.pk,
        }
        data.update(changes)
        return data

    def test_the_intake_registers_in_the_active_delegation(self):
        user = self.two_delegation_worker()
        self.client.force_login(user)
        self.switch(self.south.pk)
        form = self.client.get(reverse("core:actividades_create")).context["form"]
        self.assertNotIn("delegation", form.fields)
        self.assertIn("author_officer", form.fields)
        officers = set(form.fields["author_officer"].queryset.values_list("delegation_id", flat=True))
        self.assertEqual(officers, {self.south.pk})
        self.assertIn("Delegación Sur", form.note)

        south_officer = Officer.objects.filter(delegation=self.south, deleted_at__isnull=True).first()
        data = self.intake_data(author_officer=south_officer.pk)
        response = self.client.post(reverse("core:actividades_create"), data)
        self.assertEqual(response.status_code, 302, getattr(response, "context", None) and response.context["form"].errors)
        saved = Activity.objects.get(contact_name="Rosa Prueba")
        self.assertEqual(saved.delegation, self.south)
        self.assertEqual(saved.author_officer, south_officer)

    def test_the_intake_in_the_own_delegation_still_fills_the_responsible_alone(self):
        user = self.two_delegation_worker()
        self.client.force_login(user)
        form = self.client.get(reverse("core:actividades_create")).context["form"]
        self.assertNotIn("delegation", form.fields)
        self.assertNotIn("author_officer", form.fields)
        response = self.client.post(reverse("core:actividades_create"), self.intake_data())
        self.assertEqual(response.status_code, 302)
        saved = Activity.objects.get(contact_name="Rosa Prueba")
        self.assertEqual((saved.delegation, saved.author_officer), (self.north, user.profile.officer))

    def test_a_responsible_from_another_delegation_is_rejected_by_the_server(self):
        user = self.two_delegation_worker()
        self.client.force_login(user)
        self.switch(self.south.pk)
        north_officer = Officer.objects.filter(delegation=self.north, deleted_at__isnull=True).first()
        response = self.client.post(reverse("core:actividades_create"), self.intake_data(author_officer=north_officer.pk))
        self.assertEqual(response.status_code, 200)
        self.assertIn("author_officer", response.context["form"].errors)
        self.assertFalse(Activity.objects.filter(contact_name="Rosa Prueba").exists())

    def test_a_forged_delegation_in_the_data_is_ignored(self):
        user = self.two_delegation_worker()
        self.client.force_login(user)
        south_officer = Officer.objects.filter(delegation=self.south, deleted_at__isnull=True).first()
        data = self.intake_data(author_officer=south_officer.pk, delegation=self.south.pk)
        response = self.client.post(reverse("core:actividades_create"), data)
        saved = Activity.objects.get(contact_name="Rosa Prueba")
        self.assertEqual(response.status_code, 302)
        self.assertEqual(saved.delegation, self.north)

    def test_the_commitment_form_limits_its_lists_and_forces_the_delegation(self):
        user = self.make_user("agenda.dos", ["Funcionario", "Coordinación"], self.officer("F-1004"))
        self.client.force_login(user)
        self.switch(self.south.pk)
        form = self.client.get(reverse("core:agenda_create")).context["form"]
        self.assertEqual({t.delegation_id for t in form.fields["territory"].queryset}, {self.south.pk})
        self.assertEqual({o.delegation_id for o in form.fields["responsible_officer"].queryset}, {self.south.pk})
        self.assertEqual(list(form.fields["delegation"].queryset), [self.south])
        self.assertEqual(form.initial["delegation"], self.south.pk)
        self.assertNotEqual(form.initial.get("responsible_officer"), user.profile.officer_id)

        response = self.client.post(reverse("core:agenda_create"), self.commitment_data(delegation=self.north.pk))
        self.assertEqual(response.status_code, 302)
        saved = Commitment.objects.get(requester="Junta de vecinos")
        self.assertEqual(saved.delegation, self.south)

    def test_changing_records_is_allowed_only_in_the_active_delegation(self):
        user = self.two_delegation_worker()
        self.client.force_login(user)
        north_activity = Activity.objects.filter(delegation=self.north, deleted_at__isnull=True).first()
        south_activity = Activity.objects.filter(delegation=self.south, deleted_at__isnull=True).first()
        self.switch(self.south.pk)
        rows = {row.obj.pk: row for row in self.client.get(reverse("core:atencion")).context["rows"]}
        self.assertNotIn(north_activity.pk, rows)
        self.assertTrue(rows[south_activity.pk].edit_url)
        self.assertEqual(self.client.get(reverse("core:actividades_update", args=[south_activity.pk])).status_code, 200)
        self.assertEqual(self.client.get(reverse("core:actividades_update", args=[north_activity.pk])).status_code, 404)

    def test_a_user_without_the_delegation_gets_a_404_not_a_403_for_foreign_records(self):
        self.login_as("matias.cavieres")
        south_activity = Activity.objects.filter(delegation=self.south, deleted_at__isnull=True).first()
        self.assertEqual(self.client.get(reverse("core:actividades_update", args=[south_activity.pk])).status_code, 404)

    def test_the_forms_warn_about_unsaved_changes_before_switching(self):
        self.client.force_login(self.two_delegation_worker())
        for name in ("core:actividades_create", "core:agenda_create", "core:profile"):
            with self.subTest(name=name):
                html = self.client.get(reverse(name)).content.decode()
                self.assertIn("data-unsaved-guard", html)
                self.assertIn("core/js/context.js", html)
                self.assertIn("data-context-select", html)
        response = self.client.post(reverse("core:actividades_create"), {"requester": ""})
        self.assertContains(response, 'data-unsaved-guard="dirty"')

    def test_without_javascript_the_selector_keeps_its_button(self):
        self.login_as("supervision1")
        header = self.header()
        self.assertIn('class="btn btn-outline-primary btn-sm context-apply"', header)
        self.assertIn('name="next"', header)
        self.assertIn('name="csrfmiddlewaretoken"', header)

    def test_the_selector_is_labelled_for_screen_readers(self):
        self.login_as("supervision1")
        header = self.header()
        self.assertIn('<label class="topbar-context-label" for="contexto-delegacion">Delegación</label>', header)
        self.assertIn('id="contexto-delegacion"', header)


class UserDelegationsFormTests(ContextTestCase):

    def user_data(self, **changes):
        data = {
            "username": "ana.torres",
            "first_name": "Ana",
            "last_name": "Torres",
            "email": "ana.torres@dominio.cl",
            "groups": [Group.objects.get(name="Funcionario").pk],
            "officer": self.officer("F-1004").pk,
            "is_active": "on",
            "password1": "ClaveSegura#2026",
            "password2": "ClaveSegura#2026",
        }
        data.update(changes)
        return data

    def test_the_form_offers_the_active_delegations_with_their_codes(self):
        self.login_as("admin")
        form = self.client.get(reverse("core:usuarios_create")).context["form"]
        labels = [label for _value, label in form.fields["extra_delegations"].choices if _value]
        self.assertEqual(
            labels,
            [f"DEL{d.pk:02d} · {d.name}" for d in Delegation.objects.order_by("pk")],
        )
        html = self.client.get(reverse("core:usuarios_create")).content.decode()
        self.assertIn('data-label="Delegaciones adicionales"', html)
        self.assertIn("color-swatch", html)

    def test_creating_a_user_saves_one_row_per_role_and_delegation_and_audits_it(self):
        self.login_as("admin")
        funcionario, verificador = Group.objects.get(name="Funcionario"), Group.objects.get(name="Verificador")
        data = self.user_data(groups=[funcionario.pk, verificador.pk], extra_delegations=[self.south.pk, self.center.pk])
        response = self.client.post(reverse("core:usuarios_create"), data)
        self.assertEqual(response.status_code, 302, getattr(response, "context", None) and response.context["form"].errors)
        user = User.objects.get(username="ana.torres")
        rows = set(UserRoleDelegation.objects.filter(user=user).values_list("role_id", "delegation_id"))
        self.assertEqual(
            rows,
            {(role.pk, d.pk) for role in (funcionario, verificador) for d in (self.south, self.center)},
        )
        log = AuditLog.objects.get(event="Delegaciones autorizadas modificadas")
        self.assertEqual(log.record_id, user.pk)
        self.assertEqual(log.previous_value, "Delegaciones adicionales: (ninguna)")
        self.assertIn(f"DEL{self.south.pk:02d} · Delegación Sur", log.new_value)
        self.assertIn("Delegación Centro", log.new_value)

        self.client.logout()
        self.client.force_login(user)
        values = [value for value, _ in self.select_options(self.header())]
        self.assertEqual(values, [str(self.north.pk), str(self.south.pk), str(self.center.pk)])

    def test_editing_changes_the_rows_and_keeps_the_audit_of_each_change(self):
        self.login_as("admin")
        user = User.objects.get(username="supervision1")
        url = reverse("core:usuarios_update", args=[user.pk])
        form = self.client.get(url).context["form"]
        self.assertEqual(set(form.initial["extra_delegations"]), {self.south.pk})

        data = self.user_data(
            username="supervision1",
            first_name="Supervisión",
            last_name="Zonal",
            email="supervision1@sgr.example",
            groups=[Group.objects.get(name="Jefatura").pk],
            officer=user.profile.officer_id,
            password1="",
            password2="",
            extra_delegations=[self.center.pk],
        )
        response = self.client.post(url, data)
        self.assertEqual(response.status_code, 302, getattr(response, "context", None) and response.context["form"].errors)
        self.assertEqual([d.pk for d in extra_delegations(user)], [self.center.pk])
        log = AuditLog.objects.get(event="Delegaciones autorizadas modificadas")
        self.assertIn("Delegación Sur", log.previous_value)
        self.assertIn("Delegación Centro", log.new_value)

        data["extra_delegations"] = []
        self.client.post(url, data)
        self.assertEqual(list(extra_delegations(user)), [])
        self.assertEqual(AuditLog.objects.filter(event="Delegaciones autorizadas modificadas").count(), 2)
        self.client.post(url, data)
        self.assertEqual(AuditLog.objects.filter(event="Delegaciones autorizadas modificadas").count(), 2)

    def test_changing_the_roles_keeps_the_extra_delegations_for_the_new_roles(self):
        self.login_as("admin")
        user = User.objects.get(username="supervision1")
        url = reverse("core:usuarios_update", args=[user.pk])
        data = self.user_data(
            username="supervision1",
            first_name="Supervisión",
            last_name="Zonal",
            email="supervision1@sgr.example",
            groups=[Group.objects.get(name="Funcionario").pk],
            officer=user.profile.officer_id,
            password1="",
            password2="",
            extra_delegations=[self.south.pk],
        )
        self.assertEqual(self.client.post(url, data).status_code, 302)
        rows = UserRoleDelegation.objects.filter(user=user)
        self.assertEqual({(r.role.name, r.delegation_id) for r in rows}, {("Funcionario", self.south.pk)})
        self.assertEqual([d.pk for d in extra_delegations(user)], [self.south.pk])

    def test_a_user_that_sees_everything_cannot_have_extra_delegations(self):
        self.login_as("admin")
        coordination = Group.objects.get(name="Coordinación")
        data = self.user_data(groups=[coordination.pk], officer="", extra_delegations=[self.south.pk])
        response = self.client.post(reverse("core:usuarios_create"), data)
        self.assertEqual(response.status_code, 200)
        self.assertIn("extra_delegations", response.context["form"].errors)
        self.assertFalse(User.objects.filter(username="ana.torres").exists())

        data = self.user_data(is_superuser="on", groups=[], officer="", extra_delegations=[self.south.pk])
        response = self.client.post(reverse("core:usuarios_create"), data)
        self.assertIn("extra_delegations", response.context["form"].errors)

    def test_the_officers_own_delegation_is_not_saved_twice(self):
        self.login_as("admin")
        data = self.user_data(extra_delegations=[self.north.pk, self.south.pk])
        self.assertEqual(self.client.post(reverse("core:usuarios_create"), data).status_code, 302)
        user = User.objects.get(username="ana.torres")
        self.assertEqual([d.pk for d in extra_delegations(user)], [self.south.pk])

    def test_a_forged_delegation_id_in_the_form_is_rejected(self):
        self.login_as("admin")
        for forged in ("9999", "abc"):
            data = self.user_data(extra_delegations=[forged])
            response = self.client.post(reverse("core:usuarios_create"), data)
            self.assertEqual(response.status_code, 200)
            self.assertIn("extra_delegations", response.context["form"].errors)
        Delegation.objects.filter(pk=self.center.pk).update(status=Delegation.Status.INACTIVE)
        response = self.client.post(reverse("core:usuarios_create"), self.user_data(extra_delegations=[self.center.pk]))
        self.assertIn("extra_delegations", response.context["form"].errors)

    def test_only_who_can_change_users_can_change_the_authorized_delegations(self):
        self.login_as("jefatura1")
        self.assertEqual(self.client.get(reverse("core:usuarios_create")).status_code, 403)
        response = self.client.post(reverse("core:usuarios_create"), self.user_data(extra_delegations=[self.south.pk]))
        self.assertEqual(response.status_code, 403)
        self.assertFalse(UserRoleDelegation.objects.filter(user__username="ana.torres").exists())

    def test_the_list_and_the_detail_show_the_authorized_delegations(self):
        self.login_as("admin")
        rows = self.client.get(reverse("core:usuarios_list"), {"q": "supervision1"}).context["rows"]
        self.assertEqual([row.cells[4].text for row in rows], ["Delegación Norte y 1 más"])
        user = User.objects.get(username="supervision1")
        detail = self.client.get(reverse("core:usuarios_detail", args=[user.pk]))
        pairs = {label: cell.text for label, cell in detail.context["extra"]} if "extra" in detail.context else {}
        if pairs:
            self.assertEqual(
                pairs["Delegaciones autorizadas"],
                f"DEL{self.north.pk:02d} · Delegación Norte, DEL{self.south.pk:02d} · Delegación Sur",
            )
        self.assertContains(detail, "Delegaciones autorizadas")

    def test_the_user_list_does_not_query_per_user(self):
        self.login_as("admin")
        for index in range(8):
            user = self.make_user(f"extra{index}", ["Funcionario", "Verificador"])
            sync_extra_delegations(user, [self.south, self.center])
        with self.assertNumQueries(self.count_queries(reverse("core:usuarios_list"))):
            self.client.get(reverse("core:usuarios_list"))
        first = self.count_queries(reverse("core:usuarios_list"), {"q": "extra0"})
        many = self.count_queries(reverse("core:usuarios_list"), {"q": "extra"})
        self.assertEqual(first, many)

    def count_queries(self, url, params=None):
        from django.db import connection
        from django.test.utils import CaptureQueriesContext

        with CaptureQueriesContext(connection) as queries:
            self.client.get(url, params or {})
        return len(queries)


class LegacyBehaviourTests(ContextTestCase):

    def test_a_global_user_with_an_officer_keeps_acting_in_the_officers_delegation_while_looking_at_all(self):
        user = self.make_user("coord.funcionario", ["Coordinación", "Funcionario"], self.officer("F-1002"))
        self.client.force_login(user)
        north_activity = self.activity("María Pérez")
        south_activity = self.activity("Ana Rojas")
        rows = {row.obj.pk: row for row in self.client.get(reverse("core:atencion")).context["rows"]}
        self.assertTrue(rows[north_activity.pk].edit_url)
        self.assertFalse(rows[south_activity.pk].edit_url)
        self.assertEqual(self.client.get(reverse("core:actividades_update", args=[south_activity.pk])).status_code, 403)
        self.switch(self.south.pk)
        rows = {row.obj.pk: row for row in self.client.get(reverse("core:atencion")).context["rows"]}
        self.assertTrue(rows[south_activity.pk].edit_url)
        self.assertNotIn(north_activity.pk, rows)

    def test_a_user_without_delegation_sees_the_notice_and_no_selector(self):
        user = self.make_user("sin.delegacion", ["Funcionario"])
        self.client.force_login(user)
        response = self.client.get(reverse("core:dashboard"))
        self.assertTrue(response.context["no_delegation"])
        self.assertEqual(response.context["me_scope"]["name"], "Sin delegación asignada")
        self.assertNotContains(response, "data-context-select")
        self.assertEqual(self.switch(self.south.pk).status_code, 403)

    def test_a_user_with_only_extra_rows_and_no_officer_starts_in_the_first_one(self):
        user = self.make_user("solo.extra", ["Funcionario"])
        sync_extra_delegations(user, [self.south, self.center])
        self.client.force_login(user)
        scope = self.client.get(reverse("core:dashboard")).context["me_scope"]
        self.assertEqual(scope["name"], "Delegación Sur")
        self.assertTrue(scope["choose"])

    def test_the_django_admin_keeps_working_with_the_officers_delegation(self):
        user = self.two_delegation_worker()
        user.is_staff = True
        user.save()
        self.client.force_login(user)
        self.switch(self.south.pk)
        response = self.client.get("/admin/operations/activity/")
        self.assertEqual(response.status_code, 200)
        contacts = {activity.contact_name for activity in response.context["cl"].result_list}
        self.assertEqual(contacts, self.contacts_of(self.north))

    def test_the_new_table_is_in_the_django_admin(self):
        self.login_as("admin")
        self.assertEqual(self.client.get("/admin/accounts/userroledelegation/").status_code, 200)


class GraphicNormsTests(ContextTestCase):

    def css(self):
        return Path(finders.find("core/css/sgr.css")).read_text(encoding="utf-8")

    def test_the_whole_system_uses_arial(self):
        css = self.css()
        self.assertIn("--font-sans: Arial,", css)
        self.assertIn("--bs-body-font-family: var(--font-sans);", css)
        self.assertNotIn("Source Sans", css)
        self.assertNotIn("Source Serif", css)
        self.assertNotIn("@font-face", css)
        self.assertNotIn("Georgia", css)
        self.assertIn("--font-title: var(--font-sans);", css)
        self.assertIn(".swal2-title {\n  font-family: var(--font-title);", css)
        self.assertIn(".swal2-popup {\n  font-family: var(--bs-body-font-family);", css)

    def test_no_font_file_is_downloaded_or_preloaded(self):
        for url in (reverse("core:login"), reverse("accounts:recover")):
            with self.subTest(url=url):
                html = self.client.get(url).content.decode()
                self.assertNotIn(".woff2", html)
                self.assertNotIn('rel="preload"', html)
        self.assertIsNone(finders.find("core/vendor/fonts/SourceSans3VF-Upright.woff2"))

    def test_the_error_page_uses_arial_too(self):
        html = Path(finders.find("core/css/sgr.css")).parent.parent.parent.parent.parent.joinpath("templates/500.html").read_text(encoding="utf-8")
        self.assertIn("Arial", html)
        self.assertNotIn("Georgia", html)

    def test_the_institutional_color_is_the_one_of_the_browser_and_the_theme(self):
        css = self.css()
        self.assertIn("--bs-primary: #920000;", css)
        self.assertIn("--wine: #920000;", css)
        for url in (reverse("core:login"), reverse("accounts:recover")):
            self.assertContains(self.client.get(url), '<meta name="theme-color" content="#920000">')

    def test_the_official_arms_file_is_the_one_given_and_is_not_redrawn(self):
        path = Path(finders.find("core/img/escudo-la-serena.png"))
        data = path.read_bytes()
        self.assertTrue(data.startswith(b"\x89PNG\r\n\x1a\n"))
        width, height = struct.unpack(">II", data[16:24])
        self.assertEqual((width, height), (554, 554))

    def test_the_arms_appear_on_the_login_the_recovery_the_public_pages_and_the_header(self):
        login = self.client.get(reverse("core:login"))
        self.assertContains(login, '<img class="escudo escudo-login"')
        self.assertContains(login, "Escudo de la Ilustre Municipalidad de La Serena")
        self.assertContains(login, "core/img/escudo-la-serena.png")
        self.assertContains(self.client.get(reverse("accounts:recover")), '<img class="escudo escudo-login"')
        self.assertContains(self.client.get("/no-existe-esta-pagina/"), '<img class="escudo escudo-public"', status_code=404)
        self.login_as("matias.cavieres")
        for url in (reverse("core:dashboard"), reverse("core:atencion"), reverse("core:actividades_create")):
            self.assertContains(self.client.get(url), '<img class="escudo escudo-topbar"')

    def test_without_the_arms_file_nothing_is_drawn_in_its_place(self):
        with mock.patch("core.templatetags.core_ui.finders.find", return_value=None):
            for url in (reverse("core:login"), "/no-existe-esta-pagina/", reverse("accounts:recover")):
                response = self.client.get(url)
                self.assertNotContains(response, 'class="escudo', status_code=response.status_code)
            self.login_as("matias.cavieres")
            self.assertNotContains(self.client.get(reverse("core:dashboard")), 'class="escudo')

    def test_the_arms_keep_their_proportions_and_free_space(self):
        css = self.css()
        block = css[css.index(".escudo {") : css.index("}", css.index(".escudo {"))]
        self.assertIn("width: auto;", block)
        self.assertIn("object-fit: contain;", block)
        self.assertIn("min-height:", block)
        self.assertNotIn("filter", css[css.index(".escudo {") :][:1200])

    def test_the_institution_name_stays_in_the_header_and_the_footer(self):
        self.login_as("matias.cavieres")
        html = self.client.get(reverse("core:dashboard")).content.decode()
        self.assertIn("Ilustre Municipalidad de La Serena", html)
